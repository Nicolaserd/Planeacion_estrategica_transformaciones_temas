"""Estilo común de los Excel: encabezado 1F4E78, filas alternas, bordes, filtros y paneles inmovilizados."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL, AZUL_CLARO, AMBAR = "1F4E78", "DDEBF7", "FFE699"
BORDE = Side(style="thin", color="BFBFBF")
ESCALA_CONF = dict(start_type="num", start_value=0.5, start_color="F8696B", mid_type="num", mid_value=0.75,
                   mid_color="FFEB84", end_type="num", end_value=1, end_color="63BE7B")


def titulo(ws, filas: list[str], ancho_cols: int):
    """Filas de título combinadas sobre la tabla (la primera en negrita y grande)."""
    for i, texto in enumerate(filas, start=1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=ancho_cols)
        c = ws.cell(row=i, column=1, value=texto)
        c.font = Font(bold=(i == 1), size=14 if i == 1 else 10, color="000000" if i == 1 else "595959")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 24


def tabla(ws, fila_enc: int, encabezados: list[str], filas: list[list], anchos: list[int], pct_cols=(), wrap_cols=(),
          formatos: dict | None = None) -> int:
    """Escribe una tabla con estilo desde `fila_enc`. Devuelve la última fila escrita."""
    formatos = formatos or {}
    for j, h in enumerate(encabezados, start=1):
        c = ws.cell(row=fila_enc, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(top=BORDE, bottom=BORDE, left=BORDE, right=BORDE)
    for i, fila in enumerate(filas, start=fila_enc + 1):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = Border(top=BORDE, bottom=BORDE, left=BORDE, right=BORDE)
            c.alignment = Alignment(wrap_text=j in wrap_cols, vertical="top")
            if (i - fila_enc) % 2 == 0:
                c.fill = PatternFill("solid", fgColor=AZUL_CLARO)
            if j in pct_cols:
                c.number_format = "0.0%"
            elif j in formatos:
                c.number_format = formatos[j]
    for j, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=fila_enc + 1, column=1)
    ultima = fila_enc + max(len(filas), 1)
    ws.auto_filter.ref = f"A{fila_enc}:{get_column_letter(len(encabezados))}{ultima}"
    return ultima

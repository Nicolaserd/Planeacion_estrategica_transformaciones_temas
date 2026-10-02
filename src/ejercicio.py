"""Flujo general por ejercicio: archivo + transformación + unidad regional + fechas.

Uso:
  python src/ejercicio.py preparar  --archivo <xlsx> --transformacion "UC para la Vida" --unidad "Ubaté" \
                                     --fechas 2026-08-18 2026-08-19 [--incluir-ids 408 493]
  (Claude lee los comentarios y escribe config/ejercicios/<slug>/temas.yaml y logs/<slug>/codificacion.jsonl)
  python src/ejercicio.py validar    --ejercicio <slug>
  python src/ejercicio.py clasificar --ejercicio <slug>
  python src/ejercicio.py exportar   --ejercicio <slug>

Los temas se definen con los comentarios del ejercicio y se anclan a un componente (o concepto de "Qué busca")
de la transformación del formulario en contexto/Transformaciones_estrategicas_5_frentes.md.
"""
import argparse
import asyncio
import hashlib
import itertools
import json
import math
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
import metricas as m  # noqa: E402
from excel_estilo import ESCALA_CONF, tabla, titulo  # noqa: E402
from f00_ficha_entorno import slug  # noqa: E402
from f01_marco import bloque, extraer_secciones  # noqa: E402
from f03_limpieza_anonimizacion import anonimizar, motivo_exclusion, normalizar, tokens_nombres  # noqa: E402

DOC_MARCO = RAIZ / "contexto" / "Transformaciones_estrategicas_5_frentes.md"
FUERA = "fuera_de_la_transformacion"
GENERAL = "transformacion_en_general"   # el comentario se refiere a la transformación como un todo
ETIQUETA_ANCLA = {FUERA: "Fuera de la transformación", GENERAL: "La transformación en general"}
SIN_CONTENIDO = ("¿El `comentario` carece de contenido evaluable, por ejemplo una respuesta vacía, un 'nada', "
                 "un 'todo bien' sin más detalle o un texto sin sentido?")
TAU_POR_DEFECTO = 0.7   # params.yaml > umbral_tema_ejercicio (ver la justificación allí)
MAX_TEMAS = 10
MIN_COMENTARIOS_TEMA = 2        # piso de los temas sobre la transformación en general y de "Otros asuntos fuera..."
FRACCION_MIN_TEMA = 0.05        # los demás temas: 3 comentarios o el 5 % de N_s, lo que sea mayor
FRACCION_FUERA_PROPIO = 0.10    # un asunto fuera de la transformación tiene tema propio solo si reúne el 10 % de N_s
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
CAMPOS_TEMA = ("id", "nombre", "componente", "justificacion", "critico", "instruccion", "criterio_true",
               "criterio_false", "ejemplos_si")


# ---------------------------------------------------------------------------------------------
# Utilidades puras (probadas en tests/test_ejercicio.py)
# ---------------------------------------------------------------------------------------------

def resolver_transformacion(nombre: str, ruta: Path = DOC_MARCO) -> dict:
    """Busca la transformación en el documento marco (sin importar tildes ni mayúsculas)."""
    secciones = extraer_secciones(ruta.read_text(encoding="utf-8").splitlines())
    buscado = slug(nombre)
    s = next((s for s in secciones if slug(s["nombre"]) == buscado), None)
    if s is None:
        raise SystemExit(f"La transformación '{nombre}' no está en {ruta.name}. "
                         f"Disponibles: {[x['nombre'] for x in secciones]}")
    return {
        "id": f"T{s['numero']}",
        "nombre_oficial": s["nombre"],
        "referencia": f"{ruta.name} · '## {s['numero']}. {s['nombre']}' · líneas {s['linea_ini']}-{s['linea_fin']}",
        "enunciado": " ".join(bloque(s["cuerpo"], "Enunciado estratégico")),
        "que_busca": " ".join(bloque(s["cuerpo"], "Qué busca")),
        "componentes": [x.lstrip("- ").strip() for x in bloque(s["cuerpo"], "Componentes preliminares")],
        "texto_seccion": "\n".join(s["cuerpo"]),
    }


def componente_valido(componente: str, transformacion: dict) -> bool:
    """Un componente es válido si es uno de los componentes, una frase literal de la sección, GENERAL o FUERA."""
    return (componente in (FUERA, GENERAL) or componente in transformacion["componentes"]
            or (len(componente) >= 8 and componente in transformacion["texto_seccion"]))


def minimo_comentarios(N_s: int) -> int:
    """Mínimo de comentarios de un tema anclado a la transformación: 3 o el 5 % de N_s, lo que sea mayor."""
    return max(3, math.ceil(FRACCION_MIN_TEMA * N_s))


def componentes_del_tema(t: dict) -> list[str]:
    """Componente principal más los de los temas fusionados en él (campo opcional `otros_componentes`)."""
    return [t["componente"]] + list(t.get("otros_componentes", []))


def _ascii(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()


def parte_fechas(fechas: list[str]) -> str:
    fechas = sorted(fechas)
    return fechas[0] if len(fechas) == 1 else f"{fechas[0]}_a_{fechas[-1]}"


def slug_ejercicio(unidad: str, fechas: list[str], transformacion: str) -> str:
    return f"{slug(unidad)}_{parte_fechas(fechas)}_{slug(transformacion)}"


def nombre_archivo(prefijo: str, unidad: str, fechas: list[str], transformacion: str) -> str:
    limpio = lambda t: re.sub(r"[^A-Za-z0-9]+", "_", _ascii(t)).strip("_")  # noqa: E731
    return f"{prefijo}_{limpio(unidad)}_{parte_fechas(fechas)}_{limpio(transformacion)}.xlsx"


def nombre_hoja(unidad: str, fechas: list[str]) -> str:
    """Nombre de hoja legible y válido para Excel (máx. 31 caracteres, sin []:*?/\\)."""
    d = sorted(date.fromisoformat(f) for f in fechas)
    if len(d) == 1:
        txt = f"{d[0].day} {MESES[d[0].month - 1]} {d[0].year}"
    elif d[0].year == d[-1].year and d[0].month == d[-1].month:
        txt = f"{d[0].day}-{d[-1].day} {MESES[d[0].month - 1]} {d[0].year}"
    else:
        txt = f"{d[0].day} {MESES[d[0].month - 1]}-{d[-1].day} {MESES[d[-1].month - 1]} {d[-1].year}"
    return re.sub(r"[\[\]:*?/\\]", " ", f"{unidad} {txt}")[:31].strip()


def copia_anonimizada(ruta: Path, columna_comentario: str, com: pd.DataFrame) -> Path:
    """Copia del Excel de clasificación apta para publicar, con el mismo formato y las mismas columnas.

    Nombre pasa a "[NOMBRE]", el correo a "[CORREO]" y el comentario original a su versión anonimizada
    (`comentario_anon`). Se guarda junto al original como <nombre>_anonimizado.xlsx.
    """
    anon = dict(zip(com["id_resp"].astype(int), com["comentario_anon"]))
    wb = load_workbook(ruta)
    ws = wb.active
    enc = {ws.cell(row=1, column=j).value: j for j in range(1, ws.max_column + 1)}
    personales = {j: "[CORREO]" if slug(str(h)).startswith("correo") else "[NOMBRE]"
                  for h, j in enc.items() if h and (slug(str(h)).startswith("correo") or slug(str(h)) == "nombre")}
    j_id, j_com = enc["ID"], enc[columna_comentario]
    for i in range(2, ws.max_row + 1):
        id_resp = int(ws.cell(row=i, column=j_id).value)
        if id_resp not in anon:
            raise SystemExit(f"{ruta.name}: el ID {id_resp} no está en comentarios.parquet; no se puede anonimizar")
        for j, marca in personales.items():
            if ws.cell(row=i, column=j).value not in (None, ""):
                ws.cell(row=i, column=j).value = marca
        if ws.cell(row=i, column=j_com).value not in (None, ""):
            ws.cell(row=i, column=j_com).value = anon[id_resp] or None
    destino = ruta.with_name(f"{ruta.stem}_anonimizado.xlsx")
    wb.save(destino)
    return destino


def filtrar(df: pd.DataFrame, mapeo: dict, unidad: str, fechas: list[str], transformacion: str,
            incluir_ids=()) -> tuple[pd.DataFrame, dict]:
    """Filas del ejercicio: unidad + fechas (+ ids incluidos a mano) + transformación del formulario."""
    col_u, col_f, col_t, col_id = mapeo["unidad_regional"], mapeo["fecha"], mapeo["transformacion_declarada"], mapeo["id_resp"]
    es_unidad = df[col_u].fillna("").map(slug).str.contains(slug(unidad), regex=False)
    if not es_unidad.any():
        raise SystemExit(f"No hay filas de la unidad '{unidad}'. Unidades: {sorted(df[col_u].dropna().unique())}")
    dia = df[col_f].dt.strftime("%Y-%m-%d")
    objetivo = slug(transformacion)
    t_slug = df[col_t].map(lambda v: slug(v) if isinstance(v, str) and v.strip() else "")
    valores = set(t_slug[es_unidad]) - {""}
    otros = valores - {objetivo}
    vacios_unidad = int((es_unidad & (t_slug == "")).sum())
    if otros and vacios_unidad:
        raise SystemExit(f"El archivo mezcla formularios ({sorted(valores)}) y hay {vacios_unidad} filas sin "
                         "transformación: hay que decidir a qué formulario pertenecen.")
    if otros and objetivo not in valores:
        raise SystemExit(f"La unidad no tiene filas de '{transformacion}'; trae: {sorted(valores)}")
    es_t = (t_slug == objetivo) | ((t_slug == "") & (not otros))
    incluir = dia.isin(fechas) | df[col_id].isin(list(incluir_ids))
    sub = df[es_unidad & es_t & incluir].copy()
    if sub.empty:
        raise SystemExit("El filtro no dejó filas: revisar unidad, fechas y transformación.")
    fuera = df[es_unidad & es_t & ~dia.isin(fechas)]
    informe = {
        "filas": len(sub),
        "unidades_que_coinciden": sorted(df.loc[es_unidad, col_u].unique()),
        "transformacion_vacia_asumida": int((sub[col_t].isna()).sum()),
        "fuera_de_fecha_no_incluidas": [{"ID": int(r[col_id]), "fecha": dia[i]} for i, r in fuera.iterrows()
                                        if r[col_id] not in set(incluir_ids)],
        "fuera_de_fecha_incluidas": [{"ID": int(r[col_id]), "fecha": dia[i]} for i, r in fuera.iterrows()
                                     if r[col_id] in set(incluir_ids)],
    }
    return sub, informe


# ---------------------------------------------------------------------------------------------
# Rutas y bitácora del ejercicio
# ---------------------------------------------------------------------------------------------

def rutas(slug_ej: str) -> dict:
    return {"config": RAIZ / "config" / "ejercicios" / slug_ej, "interim": RAIZ / "data" / "interim" / slug_ej,
            "logs": RAIZ / "logs" / slug_ej, "outputs": RAIZ / "outputs"}


def carpeta_transformacion(nombre_oficial: str) -> Path:
    """outputs/<nombre oficial de la transformación>/, p. ej. outputs/UC para la Vida/ (se crea si no existe)."""
    carpeta = RAIZ / "outputs" / re.sub(r'[<>:"/\\|?*]', " ", nombre_oficial).strip()
    carpeta.mkdir(parents=True, exist_ok=True)
    return carpeta


def bitacora(slug_ej: str, texto: str) -> None:
    ruta = rutas(slug_ej)["logs"] / "decisiones.md"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    nuevo = not ruta.exists()
    with ruta.open("a", encoding="utf-8") as f:
        if nuevo:
            f.write(f"# Bitácora · {slug_ej}\n\n")
        f.write(f"- **{date.today().isoformat()}** · {texto}\n")


def cargar(slug_ej: str) -> tuple[dict, dict]:
    r = rutas(slug_ej)
    ficha = yaml.safe_load((r["config"] / "ficha.yaml").read_text(encoding="utf-8"))
    ruta_temas = r["config"] / "temas.yaml"
    temas = yaml.safe_load(ruta_temas.read_text(encoding="utf-8")) if ruta_temas.exists() else None
    return ficha, temas


# ---------------------------------------------------------------------------------------------
# Subcomandos
# ---------------------------------------------------------------------------------------------

def preparar(a):
    trans = resolver_transformacion(a.transformacion)
    fechas = sorted(pd.to_datetime(a.fechas).strftime("%Y-%m-%d"))
    slug_ej = slug_ejercicio(a.unidad, fechas, trans["nombre_oficial"])
    r = rutas(slug_ej)
    for k in ("config", "interim", "logs"):
        r[k].mkdir(parents=True, exist_ok=True)

    archivo = Path(a.archivo).resolve()
    mapeo = yaml.safe_load((RAIZ / "config" / "mapeo_columnas.yaml").read_text(encoding="utf-8"))
    df = pd.read_excel(archivo)
    col_com = mapeo["comentarios"][0]["columna"]
    faltan = [c for c in [mapeo[k] for k in ("id_resp", "grupo", "unidad_regional", "fecha", "transformacion_declarada")]
              + [col_com] if c not in df.columns]
    if faltan:
        raise SystemExit(f"Columnas esperadas que no están en el Excel: {faltan}")
    sub, informe = filtrar(df, mapeo, a.unidad, fechas, trans["nombre_oficial"], a.incluir_ids or [])
    cols_encuesta = [c for c in df.columns if df[c].notna().any()]  # sin las columnas vacías de Forms

    toks = tokens_nombres(df) if "Nombre" in df.columns else set()
    conteo, filas = {}, []
    for _, fila in sub.iterrows():
        original = fila[col_com]
        motivo = motivo_exclusion(original)
        texto = "" if motivo == "vacio" else str(original)
        anon = anonimizar(texto, conteo) if texto else ""
        palabras = {normalizar(w) for w in re.findall(r"[^\W\d_]+", anon)}
        filas.append({"id_resp": int(fila[mapeo["id_resp"]]), "id_com": f"{int(fila[mapeo['id_resp']])}-P1",
                      "grupo": fila[mapeo["grupo"]], "comentario_anon": anon, "excluido": int(motivo is not None),
                      "motivo": motivo or "", "coincide_token_nombre": int(bool(palabras & toks))})
    com = pd.DataFrame(filas)
    con = com.excluido == 0
    claves = com.loc[con, "comentario_anon"]
    mapa = {k: f"U{i + 1:03d}" for i, k in enumerate(dict.fromkeys(claves))}
    com["id_texto_unico"] = ""
    com.loc[con, "id_texto_unico"] = claves.map(mapa)

    encuesta = sub[cols_encuesta].copy()
    encuesta.insert(0, "id_com", com["id_com"].values)
    encuesta.to_parquet(r["interim"] / "encuesta.parquet", index=False)
    com.to_parquet(r["interim"] / "comentarios.parquet", index=False)

    params = yaml.safe_load((RAIZ / "config" / "params.yaml").read_text(encoding="utf-8"))
    protocolo = m.protocolo_por_tamano(int(con.sum()), **params["regla_censo"])
    ficha = {
        "slug": slug_ej,
        "archivo": str(archivo),
        "sha256_archivo": hashlib.sha256(archivo.read_bytes()).hexdigest(),
        "transformacion": {k: trans[k] for k in ("id", "nombre_oficial", "referencia")},
        "unidad": a.unidad,
        "fechas": fechas,
        "incluir_ids": list(a.incluir_ids or []),
        "columnas_encuesta": cols_encuesta,
        "columna_comentario": col_com,
        "N_filas": len(com), "N_s": int(con.sum()), "textos_unicos": len(mapa),
        "excluidos": com.loc[~con, "motivo"].value_counts().to_dict(),
        "anonimizacion": conteo,
        "protocolo": protocolo,
        "informe_filtro": informe,
        "archivos_salida": {"temas": nombre_archivo("Temas", a.unidad, fechas, trans["nombre_oficial"]),
                            "clasificacion": nombre_archivo("Clasificacion", a.unidad, fechas, trans["nombre_oficial"]),
                            "hoja": nombre_hoja(a.unidad, fechas)},
        "preparado": date.today().isoformat(),
    }
    (r["config"] / "ficha.yaml").write_text(yaml.safe_dump(ficha, allow_unicode=True, sort_keys=False, width=110),
                                             encoding="utf-8")
    bitacora(slug_ej, f"preparar: {archivo.name} · {trans['nombre_oficial']} ({trans['id']}) · {a.unidad} · {fechas} · "
                      f"incluir_ids={ficha['incluir_ids']}. N={ficha['N_filas']}, N_s={ficha['N_s']}, excluidos={ficha['excluidos']}, "
                      f"anonimización={conteo}, protocolo={protocolo['protocolo']}. Fuera de fecha no incluidas: "
                      f"{informe['fuera_de_fecha_no_incluidas']}.")

    print(json.dumps({k: ficha[k] for k in ("slug", "transformacion", "unidad", "fechas", "N_filas", "N_s", "excluidos",
                                            "anonimizacion", "protocolo", "informe_filtro", "archivos_salida")},
                     ensure_ascii=False, indent=1))
    print(f"\n=== Componentes de {trans['nombre_oficial']} (anclas válidas para los temas) ===")
    for c in trans["componentes"]:
        print(f"- {c}")
    print(f"Qué busca: {trans['que_busca']}")
    if not a.silencioso:
        print("\n=== Comentarios anonimizados para leer (N_s) ===")
        for _, x in com[con].iterrows():
            marca = " <<token de nombre>>" if x.coincide_token_nombre else ""
            print(f"[{x.id_com}]{marca} {x.comentario_anon.strip()}")


def validar(slug_ej: str, imprimir: bool = True) -> dict:
    ficha, temas = cargar(slug_ej)
    if temas is None:
        raise SystemExit(f"Falta config/ejercicios/{slug_ej}/temas.yaml")
    trans = resolver_transformacion(ficha["transformacion"]["nombre_oficial"])
    lista = temas["temas"]
    errores = []
    ids = [t.get("id") for t in lista]
    if len(ids) != len(set(ids)):
        errores.append("ids de tema repetidos")
    for t in lista:
        faltan = [k for k in CAMPOS_TEMA if k not in t]
        if faltan:
            errores.append(f"{t.get('id')}: faltan campos {faltan}")
        for cmp in componentes_del_tema(t) if "componente" in t else [""]:
            if not componente_valido(cmp, trans):
                errores.append(f"{t['id']}: el componente '{cmp}' no está en la sección de "
                               f"{trans['nombre_oficial']} del documento marco")
    if len(lista) > MAX_TEMAS:
        errores.append(f"{len(lista)} temas: el máximo es {MAX_TEMAS}; fusionar los afines")

    com = pd.read_parquet(rutas(slug_ej)["interim"] / "comentarios.parquet")
    ids_ns = set(com.loc[com.excluido == 0, "id_com"])
    ruta_cod = rutas(slug_ej)["logs"] / "codificacion.jsonl"
    cod = [json.loads(l) for l in ruta_cod.read_text(encoding="utf-8").splitlines() if l.strip()]
    leidos = [c["id_com"] for c in cod]
    if set(leidos) != ids_ns or len(leidos) != len(set(leidos)):
        errores.append(f"la codificación no cubre exactamente N_s: diferencia {sorted(ids_ns ^ set(leidos))}")
    desconocidos = {x for c in cod for x in c["temas"]} - set(ids)
    if desconocidos:
        errores.append(f"temas en la codificación que no están en temas.yaml: {desconocidos}")
    N = len(leidos)
    Y = pd.DataFrame(0, index=leidos, columns=ids)
    for c in cod:
        Y.loc[c["id_com"], [x for x in c["temas"] if x in ids]] = 1
    freq = {k: int(Y[k].sum()) for k in ids}
    minimo, umbral_fuera = minimo_comentarios(N), math.ceil(FRACCION_FUERA_PROPIO * N)
    fuera_pequenos = [t["id"] for t in lista if t["componente"] == FUERA and freq[t["id"]] < umbral_fuera]
    if len(fuera_pequenos) > 1:
        errores.append(f"asuntos fuera de la transformación con menos de {umbral_fuera} comentarios ({fuera_pequenos}): "
                       "agruparlos en un solo tema 'Otros asuntos fuera de la transformación'")
    for t in lista:
        piso = (0 if t.get("critico") else
                MIN_COMENTARIOS_TEMA if t["componente"] == GENERAL or t["id"] in fuera_pequenos else minimo)
        if freq[t["id"]] < piso:
            errores.append(f"{t['id']}: {freq[t['id']]} comentario(s), el mínimo es {piso}; fusionar con otro tema")
    pares = [(a, b, (Y[a] & Y[b]).sum() / max(1, (Y[a] | Y[b]).sum())) for a, b in itertools.combinations(ids, 2)]
    for a, b, j in pares:
        if j > 0.5:
            errores.append(f"solapamiento Jaccard {j:.2f} entre {a} y {b}: fusionar o precisar la frontera")
    if errores:
        raise SystemExit("Validación de temas fallida:\n- " + "\n- ".join(errores))
    res = {"N_s": N, "temas": len(ids), "fuera_de_la_transformacion": [t["id"] for t in lista if t["componente"] == FUERA],
           "sobre_la_transformacion_en_general": [t["id"] for t in lista if t["componente"] == GENERAL],
           "frecuencia_lectura": freq, "sin_tema": int((Y.sum(axis=1) == 0).sum()),
           "max_jaccard": round(max((j for _, _, j in pares), default=0), 2),
           "minimo_por_tema": minimo, "umbral_fuera_propio": umbral_fuera,
           "componentes_sin_tema": [c for c in trans["componentes"]
                                    if c not in {x for t in lista for x in componentes_del_tema(t)}]}
    if imprimir:
        print(json.dumps(res, ensure_ascii=False, indent=1))
    return res


def clasificar(slug_ej: str):
    from jev_cliente import Cache, cargar_clave, clasificar_preguntas, construir_preguntas_temas

    validar(slug_ej, imprimir=False)
    cargar_clave()
    ficha, temas = cargar(slug_ej)
    params = yaml.safe_load((RAIZ / "config" / "params.yaml").read_text(encoding="utf-8"))
    modelo = params["modelo_jev"]
    lista = temas["temas"]
    sin_cont = temas.get("sin_contenido", SIN_CONTENIDO)
    preguntas = construir_preguntas_temas(lista, sin_cont)
    ids_noul = [t["id"] for t in lista] + ["sin_contenido"]
    # Solo lo que se envía a Jev entra en la clave de caché (editar ejemplos o nombres no repite llamadas)
    config = {"flujo": "temas_por_ejercicio", "sin_contenido": sin_cont,
              "temas": [{k: t[k] for k in ("id", "instruccion", "criterio_true", "criterio_false")} for t in lista]}

    com = pd.read_parquet(rutas(slug_ej)["interim"] / "comentarios.parquet")
    unicos = com[com.excluido == 0].drop_duplicates("id_texto_unico")[["id_texto_unico", "comentario_anon"]].to_dict("records")
    cache = Cache()
    antes = len(cache._d)
    extra = {"slug_ejercicio": slug_ej}
    ok, errores = asyncio.run(clasificar_preguntas(unicos, preguntas, ids_noul, [], config, modelo, cache,
                                                   params["concurrencia_jev"], extra))
    if errores:
        pendientes = [u for u in unicos if u["id_texto_unico"] in {e[0] for e in errores}]
        ok2, errores = asyncio.run(clasificar_preguntas(pendientes, preguntas, ids_noul, [], config, modelo, cache, 2, extra))
        ok += ok2
    validas = {f["id_texto_unico"]: f for f in ok
               if set(f["nouls"]) == set(ids_noul) and all(0 <= v <= 1 for v in f["nouls"].values())}
    if errores or len(validas) != len(unicos):
        raise SystemExit(f"Completitud fallida: {len(validas)}/{len(unicos)} válidas; errores={errores}")

    filas = [{"id_texto_unico": k, **{f"p_{q}": v for q, v in f["nouls"].items()}, "modelo": f["modelo"],
              "input_tokens": f["input_tokens"]} for k, f in validas.items()]
    out = com.merge(pd.DataFrame(filas), on="id_texto_unico", how="left")
    out.to_parquet(rutas(slug_ej)["interim"] / "clasificacion.parquet", index=False)
    tokens = int(sum(f["input_tokens"] for f in validas.values()))
    nuevas = len(Cache()._d) - antes
    bitacora(slug_ej, f"clasificar: Jev ({', '.join(sorted({f['modelo'] for f in validas.values()}))}) · "
                      f"{len(validas)}/{len(unicos)} textos únicos válidos · {len(lista)} temas + sin_contenido · "
                      f"llamadas nuevas {nuevas} · tokens {tokens} · costo USD {tokens * params['precio_usd_por_millon_tokens'] / 1e6:.5f}.")
    print(json.dumps({"validas": f"{len(validas)}/{len(unicos)}", "llamadas_nuevas": nuevas, "tokens": tokens},
                     ensure_ascii=False))


def exportar(slug_ej: str):
    ficha, temas = cargar(slug_ej)
    trans = resolver_transformacion(ficha["transformacion"]["nombre_oficial"])
    lista = temas["temas"]
    por_id = {t["id"]: t for t in lista}
    r = rutas(slug_ej)
    cls = pd.read_parquet(r["interim"] / "clasificacion.parquet")
    enc = pd.read_parquet(r["interim"] / "encuesta.parquet")
    c = cls[cls.excluido == 0].reset_index(drop=True)
    N_s, N = len(c), len(cls)
    modelo = ", ".join(sorted(c["modelo"].dropna().unique()))
    fechas_txt = " y ".join(ficha["fechas"])
    tard = ficha["informe_filtro"]["fuera_de_fecha_incluidas"]
    extra_tard = f" (+ {len(tard)} respuesta(s) de otra fecha incluidas: ID {', '.join(str(x['ID']) for x in tard)})" if tard else ""
    rotulo = (f"{ficha['unidad']} · {fechas_txt}{extra_tard} · Transformación: {trans['nombre_oficial']} · "
              f"{N} respuestas, {N_s} con comentario · Clasificado por Jev ({modelo})")

    # Asignación por comentario
    params = yaml.safe_load((RAIZ / "config" / "params.yaml").read_text(encoding="utf-8"))
    tau = float(params.get("umbral_tema_ejercicio", TAU_POR_DEFECTO))
    ids = [t["id"] for t in lista]
    asign = {}
    for _, x in c.iterrows():
        ps = sorted(((k, float(x[f"p_{k}"])) for k in ids), key=lambda kv: -kv[1])
        asign[x.id_com] = {"asignados": [(k, p) for k, p in ps if p >= tau], "max": ps[0]}
    etiqueta = lambda k: por_id[k]["nombre"]  # noqa: E731
    comp = lambda k: ETIQUETA_ANCLA.get(por_id[k]["componente"], por_id[k]["componente"])  # noqa: E731

    def comp_con_otros(k):
        otros = [x for x in por_id[k].get("otros_componentes", []) if x not in ETIQUETA_ANCLA]
        return comp(k) + (f" (también: {'; '.join(otros)})" if otros else "")

    # ---------- Excel 1: lista de temas (1 hoja) ----------
    # Orden: componentes del documento, frases de "Qué busca", la transformación en general, fuera de la transformación
    orden_comp = trans["componentes"]
    grupo_orden = lambda c: 0 if c in orden_comp else {GENERAL: 2, FUERA: 3}.get(c, 1)  # noqa: E731
    clave_orden = lambda t: (grupo_orden(t["componente"]),  # noqa: E731
                             orden_comp.index(t["componente"]) if t["componente"] in orden_comp else 0,
                             -sum(1 for a in asign.values() if t["id"] in dict(a["asignados"])))
    filas1 = []
    orden_temas = sorted(lista, key=clave_orden)
    for t in orden_temas:
        k = t["id"]
        con_tema = [(i, p) for i, a in asign.items() for kk, p in a["asignados"] if kk == k]
        n = len(con_tema)
        lo, hi = m.wilson(n, N_s)
        mejor = max(con_tema, key=lambda ip: ip[1]) if con_tema else None
        texto_mejor = (f"({mejor[1]:.0%}) " + c.loc[c.id_com == mejor[0], "comentario_anon"].iloc[0].strip()) if mejor else ""
        filas1.append([comp_con_otros(k), t["nombre"], t["criterio_true"], n, n / N_s, lo, hi,
                       float(np.mean([p for _, p in con_tema])) if n else None, texto_mejor])
    con_tema_algun = {x for t in lista for x in componentes_del_tema(t)}
    for comp_vacio in [x for x in orden_comp if x not in con_tema_algun]:
        filas1.append([comp_vacio, "— Ningún comentario de esta unidad trató este componente", "", 0, 0.0, None, None, None, ""])
    sin_tema = sum(1 for a in asign.values() if not a["asignados"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Temas"
    titulo(ws, [f"Temas · {trans['nombre_oficial']} · {ficha['unidad']} · {fechas_txt}", rotulo,
                f"Temas definidos a partir de los comentarios y anclados a los componentes de {trans['nombre_oficial']} "
                f"({trans['referencia']}). Comentarios = cuántos comentarios clasificó Jev en el tema (confianza ≥ {tau:.0%}; un "
                f"comentario puede tener varios temas). % sobre {N_s} comentarios. Sin tema asignado: {sin_tema}. Cada respuesta "
                "es de un grupo de 5 a 8 personas. Confianza = probabilidad que asigna Jev, no exactitud."], 9)
    ult = tabla(ws, 5, ["Componente de la transformación", "Tema", "Qué incluye", "Comentarios", "% de comentarios",
                        "IC 95 % inf.", "IC 95 % sup.", "Confianza media (Jev)", "Comentario con mayor confianza (anonimizado)"],
                filas1, [30, 40, 60, 11, 11, 9, 9, 12, 80], pct_cols=(5, 6, 7, 8), wrap_cols=(1, 2, 3, 9))
    ws.conditional_formatting.add(f"H6:H{ult}", ColorScaleRule(**ESCALA_CONF))
    carpeta = carpeta_transformacion(trans["nombre_oficial"])
    ruta1 = carpeta / ficha["archivos_salida"]["temas"]
    wb.save(ruta1)

    # ---------- Excel 2: encuesta completa + clasificación de Jev (1 hoja) ----------
    # Tema principal, su componente y su %; después una columna por tema (mismo orden que el Excel 1) con el % de
    # confianza de Jev cuando asignó ese tema (≥ umbral) y vacía cuando no: así se filtra o suma cada tema en Excel.
    info = cls.set_index("id_com")
    col_com = ficha["columna_comentario"]
    cols = [x for x in ficha["columnas_encuesta"]]
    ids_orden = [t["id"] for t in orden_temas]
    encabezados = ["Rol (Tipo de actor)" if x == "Tipo de actor" else x for x in cols] + [
        "Tema principal (Jev)", "Componente de la transformación", "% confianza (Jev)"] + [etiqueta(k) for k in ids_orden]
    filas2 = []
    for _, e in enc.iterrows():
        x = info.loc[e.id_com]
        por_tema = [None] * len(ids_orden)
        if x.excluido:
            principal, componente, conf = ("(sin comentario)" if x.motivo == "vacio" else f"(sin contenido: {x.motivo})"), "", None
        else:
            a = asign[e.id_com]
            if a["asignados"]:
                k, p = a["asignados"][0]
                principal, componente, conf = etiqueta(k), comp(k), p
                for kk, pp in a["asignados"]:
                    por_tema[ids_orden.index(kk)] = pp
            else:
                principal, componente, conf = f"Sin tema asignado (más cercano: {etiqueta(a['max'][0])})", "", a["max"][1]
        valores = [None if (isinstance(v, float) and np.isnan(v)) else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
                   for v in (e[cc] for cc in cols)]
        filas2.append(valores + [principal, componente, conf] + por_tema)
    formatos = {j: "yyyy-mm-dd hh:mm" for j, cc in enumerate(cols, start=1) if str(enc[cc].dtype).startswith("datetime")}
    anchos = [70 if cc == col_com else 18 for cc in cols] + [40, 30, 11] + [18] * len(ids_orden)
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = ficha["archivos_salida"]["hoja"]
    j_com = cols.index(col_com) + 1
    j_pct = len(cols) + 3
    j_temas = list(range(j_pct + 1, j_pct + 1 + len(ids_orden)))
    ult2 = tabla(ws2, 1, encabezados, filas2, anchos, pct_cols=(j_pct, *j_temas),
                 wrap_cols=(j_com, len(cols) + 1, len(cols) + 2), formatos=formatos)
    letra = get_column_letter(j_pct)
    ws2.conditional_formatting.add(f"{letra}2:{get_column_letter(j_temas[-1])}{ult2}", ColorScaleRule(**ESCALA_CONF))
    ruta2 = carpeta / ficha["archivos_salida"]["clasificacion"]
    wb2.save(ruta2)
    ruta2_anon = copia_anonimizada(ruta2, col_com, cls)   # la que se puede publicar

    # Verificaciones
    assert len(filas2) == N == len(enc), "Excel 2: no tiene una fila por respuesta"
    assert all(f[len(cols)] for f in filas2), "Excel 2: hay filas sin 'Tema principal'"
    assert all(v is None or 0 <= v <= 1 for f in filas2 for v in f[j_pct - 1:])
    n_asig = sum(len(a["asignados"]) for a in asign.values())
    assert sum(f[3] for f in filas1) == n_asig, "Excel 1: los conteos no cuadran con las asignaciones"
    assert sum(v is not None for f in filas2 for v in f[j_pct:]) == n_asig, "Excel 2: las columnas de tema no cuadran"
    bitacora(slug_ej, f"exportar: {ruta1.name} (1 hoja, {len(lista)} temas, {n_asig} asignaciones, {sin_tema} sin tema) y "
                      f"{ruta2.name} (1 hoja '{ws2.title}', {len(filas2)} filas = respuestas, una columna por tema; incluye "
                      f"Nombre y Correo: archivo solo para uso local) y su copia publicable {ruta2_anon.name} (sin Nombre ni "
                      "Correo, comentario anonimizado).")
    print(json.dumps({"temas": ruta1.name, "clasificacion": ruta2.name, "hoja": ws2.title, "filas_excel2": len(filas2),
                      "asignaciones": n_asig, "sin_tema": sin_tema}, ensure_ascii=False, indent=1))


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = p.add_subparsers(dest="cmd", required=True)
    pp = sp.add_parser("preparar")
    pp.add_argument("--archivo", required=True)
    pp.add_argument("--transformacion", required=True)
    pp.add_argument("--unidad", required=True)
    pp.add_argument("--fechas", nargs="+", required=True)
    pp.add_argument("--incluir-ids", nargs="*", type=int, default=[])
    pp.add_argument("--silencioso", action="store_true", help="no imprimir los comentarios")
    for nombre in ("validar", "clasificar", "exportar"):
        sp.add_parser(nombre).add_argument("--ejercicio", required=True)
    a = p.parse_args()
    if a.cmd == "preparar":
        preparar(a)
    elif a.cmd == "validar":
        validar(a.ejercicio)
    elif a.cmd == "clasificar":
        clasificar(a.ejercicio)
    else:
        exportar(a.ejercicio)


if __name__ == "__main__":
    main()

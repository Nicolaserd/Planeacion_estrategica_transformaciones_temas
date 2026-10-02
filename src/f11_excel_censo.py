"""Entregables del censo: lista de temas y clasificación de Jev comentario por comentario.

Solo contiene lo que clasificó Jev (sin la lectura de Claude), por decisión de Nicolás (2026-10-01).
Genera:
  outputs/lista_temas_<slug>.xlsx
  outputs/clasificacion_jev_<slug>.xlsx
"""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
import metricas as m  # noqa: E402
from excel_estilo import AMBAR, ESCALA_CONF, tabla, titulo  # noqa: E402
from f00_ficha_entorno import bitacora, cargar_ficha  # noqa: E402
from f05_preguntas_jev import cargar_config  # noqa: E402

TIPOS = {"oportunidad_de_mejora": "Oportunidad de mejora", "reconocimiento": "Reconocimiento",
         "neutro_observacion": "Neutro / observación"}
NOTA_CONF = ("Confianza = probabilidad que asigna Jev (0 a 100 %); un tema se asigna si su confianza es de 50 % o más. "
             "No es la exactitud del acierto: aún no hay validación con etiquetado humano.")


def main():
    ficha = cargar_ficha()
    params, marco, tax, hashes = cargar_config()
    slug, tau, seed = ficha["slug_ejercicio"], params["censo"]["tau_fijo"], params["seed"]
    nombres_t = {t["id"]: t["nombre_oficial"] for t in marco}
    temas = tax["temas"]
    nombre_tema = {t["id"]: t["nombre"] for t in temas}
    padre = {t["id"]: t["transformacion_padre"] for t in temas}
    ids_t = [t["id"] for t in marco]

    df = pd.read_parquet(RAIZ / "data" / "interim" / "clasificacion.parquet")
    c = df[df.excluido == 0].copy().reset_index(drop=True)
    N = len(c)
    modelo = ", ".join(sorted(c["modelo"].unique()))
    fechas = " y ".join(ficha["fechas_desarrollo"])
    rotulo = (f"Ejercicio: Seccional Ubaté · {fechas} (+ 2 respuestas tardías, ID 408 y 493) · "
              f"N = {N} grupos con comentario · Clasificado por Jev ({modelo}) · Taxonomía v{hashes['version_taxonomia']}")
    rng = np.random.default_rng(seed)
    c["_desempate"] = rng.random(N)

    # ---------- Excel 1: lista de temas (solo Jev) ----------
    wb = Workbook()
    ws = wb.active
    ws.title = "Temas"
    filas = []
    for t in sorted(temas, key=lambda t: (t["transformacion_padre"], -int(c[f"tema_{t['id']}"].sum()))):
        k = t["id"]
        asignados = c[c[f"tema_{k}"] == 1]
        n = len(asignados)
        lo, hi = m.wilson(n, N)
        conf_media = float(asignados[f"p_tema_{k}"].mean()) if n else None
        mejora = float((asignados["tipo"] == "oportunidad_de_mejora").mean()) if n else None
        cand = asignados.sort_values([f"p_tema_{k}", "_desempate"], ascending=[False, True])
        voces = pd.concat([cand[cand.revisar == 0], cand[cand.revisar == 1]]).head(params["n_voces_por_tema"])
        txt_voces = "\n".join(f"• ({v[f'p_tema_{k}']:.0%}) {v.comentario_anon.strip()}" for _, v in voces.iterrows())
        filas.append([f"{t['transformacion_padre']} · {nombres_t.get(t['transformacion_padre'], 'Fuera del marco')}",
                      t["nombre"], t["criterio_true"], n, n / N, lo, hi, conf_media, mejora, txt_voces])
    titulo(ws, ["Temas clasificados por Jev · Seccional Ubaté · UC para la Vida", rotulo,
                "Grupos = número de grupos (de 5 a 8 personas) a cuyo comentario Jev asignó el tema. Un comentario puede "
                "tener varios temas, por eso los porcentajes suman más de 100 %. " + NOTA_CONF], 10)
    ultima = tabla(ws, 5, ["Transformación", "Tema", "Qué incluye", "Grupos", "% de N", "IC 95 % inf.", "IC 95 % sup.",
                           "Confianza media de Jev", "% oportunidad de mejora", "Comentarios con mayor confianza (anonimizados)"],
                   filas, [26, 34, 55, 8, 8, 8, 8, 11, 11, 85], pct_cols=(5, 6, 7, 8, 9), wrap_cols=(1, 2, 3, 10))
    ws.conditional_formatting.add(f"H6:H{ultima}", ColorScaleRule(**ESCALA_CONF))

    ws2 = wb.create_sheet("Por_transformacion")
    filas2 = []
    for tid in ids_t:
        n = int(c[f"trans_{tid}"].sum())
        lo, hi = m.wilson(n, N)
        conf = float(c.loc[c[f"trans_{tid}"] == 1, f"p_trans_{tid}"].mean()) if n else None
        np_ = int((c["transformacion_principal"] == tid).sum())
        n_tema = {t["id"]: int(c["tema_" + t["id"]].sum()) for t in temas if padre[t["id"]] == tid}
        top = sorted(n_tema, key=lambda k: -n_tema[k])[:5]
        filas2.append([tid, nombres_t[tid], n, n / N, lo, hi, conf, np_, np_ / N,
                       ", ".join(f"{nombre_tema[k]} ({n_tema[k]})" for k in top)])
    nn = int((c["transformacion_principal"] == "ninguna").sum())
    filas2.append(["ninguna", "No se relaciona con ninguna de las cinco", None, None, None, None, None, nn, nn / N, ""])
    titulo(ws2, ["Transformaciones según Jev · Seccional Ubaté", rotulo,
                 "C-G: la transformación aplica al comentario (confianza ≥ 50 %; un comentario puede tocar varias). "
                 "H-I: transformación PRINCIPAL que eligió Jev (suma 100 % con 'ninguna')."], 10)
    tabla(ws2, 5, ["id", "Transformación", "Grupos", "% de N", "IC 95 % inf.", "IC 95 % sup.", "Confianza media de Jev",
                   "Como principal", "% como principal", "Temas más frecuentes (grupos)"],
          filas2, [8, 32, 8, 8, 8, 8, 11, 10, 10, 80], pct_cols=(4, 5, 6, 7, 9), wrap_cols=(2, 10))

    ws3 = wb.create_sheet("Notas")
    notas = [
        ["Ficha", f"{ficha['nombre_ejercicio']} · fechas {fechas} · lugar {ficha['lugar_desarrollo']}"],
        ["Unidad de análisis", "Grupo de 5 a 8 personas (una respuesta por grupo). No se pondera por tamaño del grupo."],
        ["Clasificador", f"Jev ({modelo}), una petición por comentario anonimizado; Jev no conoce el grupo, la fecha ni el formulario."],
        ["Umbral", f"Un tema o transformación se asigna con confianza ≥ {tau:.0%}. Fijado de antemano."],
        ["Protocolo", f"Censo (todos los comentarios): {m.protocolo_por_tamano(N, **params['regla_censo'])['protocolo']}"],
        ["Limitación", "Aún no hay etiquetado humano: las cifras describen lo que asigna Jev, no su exactitud. "
                       "Jev tiende a asignar más temas de lo necesario cuando la confianza está entre 50 % y 70 %."],
        ["Versiones", f"marco sha256 {hashes['marco']} · taxonomía sha256 {hashes['taxonomia']}"],
    ]
    tabla(ws3, 1, ["Concepto", "Detalle"], notas, [22, 140], wrap_cols=(2,))
    ruta1 = RAIZ / "outputs" / f"lista_temas_{slug}.xlsx"
    wb.save(ruta1)

    # ---------- Excel 2: clasificación de Jev comentario por comentario ----------
    wb2 = Workbook()
    idx = wb2.active
    idx.title = "Indice"
    hoja_t5, hoja_largo = "T5 UC para la Vida", "Temas_por_comentario"

    filas_cls, filas_largo = [], []
    for _, r in c.iterrows():
        asig = sorted([(t["id"], r[f"p_tema_{t['id']}"]) for t in temas if r[f"tema_{t['id']}"] == 1], key=lambda x: -x[1])
        if asig:
            tp, ptp = asig[0]
            tema_principal = f"{padre[tp]} · {nombre_tema[tp]}"
        else:
            tema_principal, ptp = "(sin tema específico)", max(r[f"p_tema_{t['id']}"] for t in temas)
        otros = "; ".join(f"{padre[k]} · {nombre_tema[k]} ({p:.0%})" for k, p in asig[1:])
        princ = r["transformacion_principal"]
        tambien = "; ".join(f"{tid} ({r[f'p_trans_{tid}']:.0%})" for tid in ids_t if r[f"trans_{tid}"] == 1 and tid != princ)
        filas_cls.append([r.id_com, r.comentario_anon.strip(), tema_principal, float(ptp), otros,
                          f"{princ} · {nombres_t.get(princ, 'Ninguna')}", float(r.conf_principal), tambien,
                          TIPOS[r.tipo], float(r.conf_tipo), "Sí" if r.revisar else "No", r.motivo_revision, r.grupo])
        for rango, (k, p) in enumerate(asig, start=1):
            filas_largo.append([r.id_com, r.comentario_anon.strip(), rango, f"{padre[k]} · {nombres_t[padre[k]]}",
                                nombre_tema[k], float(p)])
        if not asig:
            filas_largo.append([r.id_com, r.comentario_anon.strip(), None, "", "(sin tema específico)", float(ptp)])
    filas_cls.sort(key=lambda f: (f[2], -f[3]))

    ws = wb2.create_sheet(hoja_t5)
    ws.sheet_properties.tabColor = "70AD47"
    titulo(ws, [f"Unidad regional: Seccional Ubaté | T5: {nombres_t['T5']}", rotulo,
                NOTA_CONF + " Filas en ámbar: conviene revisarlas."], 13)
    ultima = tabla(ws, 5, ["id_com", "Comentario (anonimizado)", "Tema principal (Jev)", "% confianza tema",
                           "Otros temas (Jev, % confianza)", "Transformación principal (Jev)", "% confianza transformación",
                           "También aporta a (Jev)", "Tipo de aporte (Jev)", "% confianza tipo", "Revisar", "Motivo", "Grupo"],
                   filas_cls, [8, 70, 38, 11, 60, 30, 12, 18, 20, 10, 8, 30, 22],
                   pct_cols=(4, 7, 10), wrap_cols=(2, 3, 5, 6, 8, 12, 13))
    for col in ("D", "G", "J"):
        ws.conditional_formatting.add(f"{col}6:{col}{ultima}", ColorScaleRule(**ESCALA_CONF))
    ws.conditional_formatting.add(f"A6:M{ultima}", FormulaRule(formula=['$K6="Sí"'], fill=PatternFill("solid", fgColor=AMBAR)))

    wl = wb2.create_sheet(hoja_largo)
    titulo(wl, ["Temas que asignó Jev a cada comentario, con su % de confianza", rotulo,
                "Una fila por cada tema asignado (confianza ≥ 50 %). Use el filtro de la columna 'Tema' para ver todos los "
                "comentarios de un tema."], 6)
    ult_l = tabla(wl, 5, ["id_com", "Comentario (anonimizado)", "Orden", "Transformación del tema", "Tema (Jev)", "% confianza"],
                  filas_largo, [8, 80, 7, 32, 45, 11], pct_cols=(6,), wrap_cols=(2, 4, 5))
    wl.conditional_formatting.add(f"F6:F{ult_l}", ColorScaleRule(**ESCALA_CONF))

    det = wb2.create_sheet("Detalle_probabilidades")
    enc = (["id_com", "comentario"] + [f"{tid} {nombres_t[tid]}" for tid in ids_t] +
           [f"principal: {k}" for k in ids_t + ["ninguna"]] + [f"{padre[t['id']]} · {t['nombre']}" for t in temas] +
           [f"tipo: {TIPOS[k]}" for k in TIPOS] + ["sin contenido"])
    filas_det = [[r.id_com, r.comentario_anon.strip()] + [r[f"p_trans_{t}"] for t in ids_t] +
                 [r.get(f"p_principal_{k}", 0.0) for k in ids_t + ["ninguna"]] + [r[f"p_tema_{t['id']}"] for t in temas] +
                 [r[f"p_tipo_{k}"] for k in TIPOS] + [r.p_sin_contenido] for _, r in c.iterrows()]
    ult = tabla(det, 1, enc, filas_det, [8, 60] + [11] * (len(enc) - 2), pct_cols=tuple(range(3, len(enc) + 1)), wrap_cols=(2,))
    det.conditional_formatting.add(f"C2:{get_column_letter(len(enc))}{ult}", ColorScaleRule(
        start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color="5B9BD5"))

    sc = wb2.create_sheet("Sin_contenido")
    excl = df[df.excluido == 1]
    filas_sc = [[r.id_com, r.grupo, "Sin comentario" if r.motivo == "vacio" else f"Excluido: {r.motivo}", r.comentario_anon]
                for _, r in excl.iterrows()]
    filas_sc += [[r.id_com, r.grupo, f"Jev: sin contenido evaluable ({r.p_sin_contenido:.0%})", r.comentario_anon.strip()]
                 for _, r in c[c.sin_contenido == 1].iterrows()]
    titulo(sc, ["Unidad regional: Seccional Ubaté | Sin_contenido", rotulo,
                "Grupos sin comentario, excluidos por el filtro de ruido o marcados por Jev como sin contenido evaluable."], 4)
    tabla(sc, 5, ["id_com", "Grupo", "Motivo", "Comentario"], filas_sc, [8, 30, 40, 40], wrap_cols=(2, 3, 4))

    titulo(idx, ["Clasificación de Jev · Seccional Ubaté · UC para la Vida", rotulo, NOTA_CONF], 6)
    tabla(idx, 5, ["Hoja", "Contenido", "Filas", "Confianza media (tema principal)", "n para revisar"], [], [26, 60, 10, 16, 12])
    q, ql = f"'{hoja_t5}'", f"'{hoja_largo}'"
    filas_idx = [
        (hoja_t5, "Un comentario por fila con su tema principal y los demás temas que asignó Jev",
         f"=COUNTA({q}!A6:A1000)", f"=IFERROR(AVERAGE({q}!D6:D1000),0)", f'=COUNTIF({q}!K6:K1000,"Sí")'),
        (hoja_largo, "Una fila por cada tema asignado, con su % de confianza (para filtrar por tema)",
         f"=COUNTA({ql}!A6:A1000)", f"=IFERROR(AVERAGE({ql}!F6:F1000),0)", None),
        ("Detalle_probabilidades", "Todas las probabilidades de Jev por comentario", "=COUNTA('Detalle_probabilidades'!A2:A1000)", None, None),
        ("Sin_contenido", "Grupos sin comentario o sin contenido evaluable", "=COUNTA('Sin_contenido'!A6:A1000)", None, None),
    ]
    for i, (hoja, desc, f_n, f_conf, f_rev) in enumerate(filas_idx, start=6):
        idx[f"A{i}"] = f'=HYPERLINK("#\'{hoja}\'!A1","{hoja}")'
        idx[f"A{i}"].font = Font(color="0563C1", underline="single")
        idx[f"B{i}"], idx[f"C{i}"] = desc, f_n
        if f_conf:
            idx[f"D{i}"] = f_conf
            idx[f"D{i}"].number_format = "0.0%"
        if f_rev:
            idx[f"E{i}"] = f_rev
    ruta2 = RAIZ / "outputs" / f"clasificacion_jev_{slug}.xlsx"
    wb2.save(ruta2)

    # Verificaciones (pandas) y bitácora
    n_asig = int(c[[f"tema_{t['id']}" for t in temas]].values.sum())
    sin_tema = sum(1 for f in filas_cls if f[2] == "(sin tema específico)")
    assert len(filas_cls) == N and len({f[0] for f in filas_cls}) == N, "hoja T5: conteo o ids repetidos"
    assert len(filas_largo) == n_asig + sin_tema, "hoja larga: no cuadra con las asignaciones"
    assert all(0 <= f[3] <= 1 for f in filas_cls) and all(0 <= f[5] <= 1 for f in filas_largo)
    bitacora(f"Entregables (solo Jev, a pedido de Nicolás): {ruta1.name} y {ruta2.name}. Hoja T5 = {N} comentarios; "
             f"hoja Temas_por_comentario = {len(filas_largo)} filas ({n_asig} asignaciones + {sin_tema} sin tema); "
             f"Sin_contenido = {len(filas_sc)}. Se retiraron del Excel las columnas de la lectura de Claude y del acuerdo Jev–Claude "
             f"(siguen registrados en esta bitácora).")
    print(json.dumps({"archivos": [ruta1.name, ruta2.name], "N": N, "asignaciones_tema": n_asig,
                      "filas_temas_por_comentario": len(filas_largo), "sin_tema": sin_tema}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()

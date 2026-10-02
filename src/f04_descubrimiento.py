"""Fase 4 · Descubrimiento de temas dentro del marco (protocolo B: censo).

La codificación abierta la hizo Claude leyendo los textos (logs/codificacion_abierta.jsonl);
este script solo cuenta, valida y genera los entregables. Nada se clasifica aquí con palabras clave.
"""
import hashlib
import itertools
import json
import sys
from pathlib import Path

import matplotlib
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
import metricas as m  # noqa: E402
from f00_ficha_entorno import bitacora, cargar_ficha  # noqa: E402

AZUL, AZUL_CLARO = "1F4E78", "DDEBF7"


def _str_multilinea(dumper, valor):
    return dumper.represent_scalar("tag:yaml.org,2002:str", valor, style="|" if "\n" in valor else None)


yaml.SafeDumper.add_representer(str, _str_multilinea)


def estilo_hoja(ws, anchos: dict):
    borde = Side(style="thin", color="BFBFBF")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, fila in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in fila:
            c.border = Border(top=borde, bottom=borde, left=borde, right=borde)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=AZUL_CLARO)
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    ficha = cargar_ficha()
    params = yaml.safe_load((RAIZ / "config" / "params.yaml").read_text(encoding="utf-8"))
    marco = yaml.safe_load((RAIZ / "config" / "marco_estrategico.yaml").read_text(encoding="utf-8"))
    nombres_t = {t["id"]: t["nombre_oficial"] for t in marco}
    red = yaml.safe_load((RAIZ / "config" / "taxonomia_redaccion.yaml").read_text(encoding="utf-8"))
    temas = red["temas"]
    ids_temas = [t["id"] for t in temas]
    padre = {t["id"]: t["transformacion_padre"] for t in temas}

    df = pd.read_parquet(RAIZ / "data" / "interim" / "comentarios.parquet")
    d = df[df.excluido == 0].set_index("id_com")
    cod = [json.loads(l) for l in (RAIZ / "logs" / "codificacion_abierta.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]

    # Validaciones
    leidos = [c["id_com"] for c in cod]
    assert len(leidos) == len(set(leidos)), "id_com repetido en la codificación"
    assert set(leidos) == set(d.index), f"La codificación no cubre exactamente N_s: {set(d.index) ^ set(leidos)}"
    desconocidos = {tm for c in cod for tm in c["temas"]} - set(ids_temas)
    assert not desconocidos, f"Temas no definidos en la taxonomía: {desconocidos}"
    assert set(padre.values()) <= {"T1", "T2", "T3", "T4", "T5", "fuera_del_marco"}

    N = len(cod)
    Y = pd.DataFrame(0, index=leidos, columns=ids_temas)
    for c in cod:
        Y.loc[c["id_com"], c["temas"]] = 1

    # Frecuencias con IC de Wilson
    freq = {}
    for tid in ids_temas:
        n = int(Y[tid].sum())
        lo, hi = m.wilson(n, N)
        freq[tid] = {"n": n, "pct": round(n / N, 4), "ic95": [round(lo, 4), round(hi, 4)]}
    pocos = [t["id"] for t in temas if freq[t["id"]]["n"] < params["censo"]["min_comentarios_tema"] and not t["critico"]]
    assert not pocos, f"Temas con menos de {params['censo']['min_comentarios_tema']} comentarios: {pocos}"

    # Solapamiento (Jaccard en D)
    pares = []
    for a, b in itertools.combinations(ids_temas, 2):
        inter = int((Y[a] & Y[b]).sum())
        union = int((Y[a] | Y[b]).sum())
        j = inter / union if union else 0.0
        pares.append({"tema_a": a, "tema_b": b, "interseccion": inter, "union": union, "jaccard": round(j, 3)})
    pares = sorted(pares, key=lambda p: -p["jaccard"])
    mayores = [p for p in pares if p["jaccard"] > 0.5]

    # Cobertura y fuera del marco (en D = censo; optimista porque no hay V disjunta)
    con_tema = int((Y.sum(axis=1) > 0).sum())
    cob = {"k": con_tema, "n": N, "pct": round(con_tema / N, 4), "ic95": [round(x, 4) for x in m.wilson(con_tema, N)]}
    temas_fuera = [t for t in ids_temas if padre[t] == "fuera_del_marco"]
    solo_fuera = int(((Y[temas_fuera].sum(axis=1) > 0) & (Y[[t for t in ids_temas if t not in temas_fuera]].sum(axis=1) == 0)).sum()) if temas_fuera else 0
    fuera = {"k": solo_fuera, "n": N, "pct": round(solo_fuera / N, 4), "ic95": [round(x, 4) for x in m.wilson(solo_fuera, N)]}

    # Transformación derivada de los temas (un comentario cuenta si tiene al menos un tema hijo)
    por_t = {}
    for tid in ["T1", "T2", "T3", "T4", "T5"]:
        hijos = [t for t in ids_temas if padre[t] == tid]
        k = int((Y[hijos].sum(axis=1) > 0).sum()) if hijos else 0
        por_t[tid] = {"k": k, "pct": round(k / N, 4), "ic95": [round(x, 4) for x in m.wilson(k, N)]}

    # Curva de saturación en el orden de lectura (lotes con semilla)
    vistos, curva, nuevos_lote = set(), [], {}
    for i, c in enumerate(cod, start=1):
        nuevos = set(c["temas"]) - vistos
        vistos |= set(c["temas"])
        curva.append(len(vistos))
        nuevos_lote.setdefault(c["lote"], set()).update(nuevos)
    nuevos_por_lote = {lote: sorted(v) for lote, v in nuevos_lote.items()}
    codigos_unicos = len({k for c in cod for k in c["codigos"]})

    slug = ficha["slug_ejercicio"]
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=150)
    ax.plot(range(1, N + 1), curva, marker="o", ms=3, color="#1F4E78")
    for lote in sorted(nuevos_lote)[1:]:
        x = sum(1 for c in cod if c["lote"] < lote)
        ax.axvline(x + 0.5, color="grey", ls="--", lw=0.8)
    ax.set_ylim(0, len(ids_temas) + 1)
    ax.set_xlim(0, N + 1)
    ax.set_xlabel("Comentarios leídos (orden aleatorio, semilla 2026)")
    ax.set_ylabel("Temas distintos acumulados")
    fig.suptitle(f"Los {len(ids_temas)} temas aparecen en los primeros {curva.index(len(ids_temas)) + 1} de {N} comentarios",
                 fontsize=11, x=0.06, ha="left")
    ax.set_title(f"Ubaté · {', '.join(ficha['fechas_desarrollo'])} · censo N = {N} grupos · lotes de 15 · "
                 "temas consolidados tras leer todo (curva optimista)", fontsize=8, color="grey", loc="left")
    fig.tight_layout()
    fig.savefig(RAIZ / "outputs" / "f04_curva_saturacion.png")
    plt.close(fig)

    # Taxonomía v1 versionada
    salida_temas = []
    for t in temas:
        t2 = dict(t)
        t2["frecuencia_D"] = {"n": freq[t["id"]]["n"], "pct": freq[t["id"]]["pct"]}
        salida_temas.append(t2)
    tax = {"version": red["version"], "protocolo": params["protocolo"], "temas": salida_temas,
           "tipo": red["tipo"], "sin_contenido": red["sin_contenido"]}
    ruta_tax = RAIZ / "config" / f"taxonomia_v{red['version']}.yaml"
    ruta_tax.write_text(yaml.safe_dump(tax, allow_unicode=True, sort_keys=False, width=110), encoding="utf-8")
    sha_tax = hashlib.sha256(ruta_tax.read_bytes()).hexdigest()
    sha_marco = hashlib.sha256((RAIZ / "config" / "marco_estrategico.yaml").read_bytes()).hexdigest()

    # Excel legible: Transformación → Tema
    wb = Workbook()
    ws = wb.active
    ws.title = "Arbol"
    ws.append(["Transformación", "Tema", "id", "n en D", "% de N", "IC 95 % inf.", "IC 95 % sup.",
               "Justificación en el marco", "Nota de frontera", "Crítico"])
    orden_t = ["T1", "T2", "T3", "T4", "T5", "fuera_del_marco"]
    for tid in orden_t:
        for t in sorted([t for t in temas if t["transformacion_padre"] == tid], key=lambda t: -freq[t["id"]]["n"]):
            f = freq[t["id"]]
            ws.append([f"{tid} · {nombres_t.get(tid, 'Fuera del marco')}", t["nombre"], t["id"], f["n"],
                       f["pct"], f["ic95"][0], f["ic95"][1], t["justificacion_marco"], t["nota_frontera"],
                       "Sí" if t["critico"] else "No"])
    for fila in ws.iter_rows(min_row=2, min_col=5, max_col=7):
        for c in fila:
            c.number_format = "0.0%"
    estilo_hoja(ws, {"A": 30, "B": 42, "C": 30, "D": 8, "E": 9, "F": 10, "G": 10, "H": 50, "I": 50, "J": 8})

    ws2 = wb.create_sheet("Preguntas_Jev")
    ws2.append(["Transformación", "Tema", "Instrucción", "Cuenta (criterio_true)", "No cuenta (criterio_false)",
                "Ejemplos sí", "Ejemplo vecino que no es"])
    for t in temas:
        ws2.append([t["transformacion_padre"], t["nombre"], t["instruccion"], t["criterio_true"], t["criterio_false"],
                    "\n".join(t["ejemplos_si"]), "\n".join(t["ejemplos_no_cercanos"])])
    estilo_hoja(ws2, {"A": 10, "B": 32, "C": 45, "D": 55, "E": 50, "F": 45, "G": 40})

    ws3 = wb.create_sheet("Codificacion")
    ws3.append(["id_com", "lote", "comentario_anon", "temas", "códigos abiertos"])
    for c in cod:
        ws3.append([c["id_com"], c["lote"], d.loc[c["id_com"], "comentario_anon"],
                    "; ".join(c["temas"]) or "(sin tema)", "; ".join(c["codigos"])])
    estilo_hoja(ws3, {"A": 9, "B": 6, "C": 80, "D": 45, "E": 50})

    ws4 = wb.create_sheet("Solapamiento")
    ws4.append(["tema_a", "tema_b", "intersección", "unión", "Jaccard"])
    for p in pares[:20]:
        ws4.append([p["tema_a"], p["tema_b"], p["interseccion"], p["union"], p["jaccard"]])
    estilo_hoja(ws4, {"A": 36, "B": 36, "C": 12, "D": 8, "E": 9})
    wb.save(RAIZ / "outputs" / "f04_taxonomia_v1.xlsx")

    # Resumen en Markdown
    lin = [f"# Fase 4 · Resumen del descubrimiento de temas",
           "", f"**Ejercicio:** Ubaté · {', '.join(ficha['fechas_desarrollo'])} (+ ID 408 y 493) · `{slug}`",
           f"**Protocolo:** B, censo. D = los {N} comentarios con contenido (grupos de 5 a 8 personas).",
           f"**Taxonomía v{red['version']}:** sha256 `{sha_tax}` · **Marco:** sha256 `{sha_marco}`", "",
           "## Árbol Transformación → Tema", "",
           "| Transformación | Tema | n | % (IC 95 %) |", "|---|---|---|---|"]
    for tid in orden_t:
        for t in sorted([t for t in temas if t["transformacion_padre"] == tid], key=lambda t: -freq[t["id"]]["n"]):
            f = freq[t["id"]]
            lin.append(f"| {tid} {nombres_t.get(tid, 'Fuera del marco')} | {t['nombre']} | {f['n']} | "
                       f"{100 * f['pct']:.1f} % ({100 * f['ic95'][0]:.1f}–{100 * f['ic95'][1]:.1f}) |")
    lin += ["", "## Transformación derivada de los temas (codificación de Claude, no de Jev)", "",
            "| Transformación | comentarios con ≥ 1 tema hijo | % (IC 95 %) |", "|---|---|---|"]
    lin += [f"| {tid} {nombres_t[tid]} | {v['k']}/{N} | {100 * v['pct']:.1f} % ({100 * v['ic95'][0]:.1f}–{100 * v['ic95'][1]:.1f}) |"
            for tid, v in por_t.items()]
    lin += ["", "Multi-etiqueta: los porcentajes suman más de 100 %."]
    lin += ["", "## Saturación", "",
            f"- Códigos abiertos distintos: {codigos_unicos}; temas consolidados: {len(ids_temas)}.",
            *[f"- Lote {lote}: {len(v)} temas nuevos" + (f" ({', '.join(v)})" if v else "") for lote, v in nuevos_por_lote.items()],
            f"- Criterio del prompt (≥ 300 leídos y 2 lotes seguidos sin temas nuevos): **no evaluable** con N = {N}.",
            "", "## Cobertura y fuera del marco", "",
            f"- Comentarios con al menos un tema: {cob['k']}/{N} = {100 * cob['pct']:.1f} % (IC {100 * cob['ic95'][0]:.1f}–{100 * cob['ic95'][1]:.1f}). "
            "Es una cobertura **dentro de la muestra**: no hay comentarios disjuntos para la validación V (4.6), así que es optimista.",
            f"- Solo con temas fuera del marco: {fuera['k']}/{N}.",
            "", "## Solapamiento", "",
            f"- Pares con Jaccard > 0,5: {len(mayores)}." + ("" if not mayores else " " + "; ".join(f"{p['tema_a']}–{p['tema_b']} ({p['jaccard']})" for p in mayores)),
            f"- Mayor Jaccard: {pares[0]['tema_a']}–{pares[0]['tema_b']} = {pares[0]['jaccard']}."]
    (RAIZ / "outputs" / "f04_resumen_descubrimiento.md").write_text("\n".join(lin) + "\n", encoding="utf-8")

    bitacora(f"Fase 4: {N} comentarios codificados (censo), {codigos_unicos} códigos abiertos -> {len(ids_temas)} temas; "
             f"cobertura en D {cob['k']}/{N}; Jaccard>0,5: {len(mayores)}; taxonomía v{red['version']} borrador sha256={sha_tax[:16]}…; "
             f"marco enriquecido (4.4) sha256={sha_marco[:16]}… (pendientes de aprobación PC4).")
    print(json.dumps({"N": N, "temas": len(ids_temas), "codigos_unicos": codigos_unicos, "por_transformacion": por_t,
                      "cobertura": cob, "fuera_del_marco": fuera, "jaccard_mayor_0_5": mayores, "top5_jaccard": pares[:5],
                      "nuevos_por_lote": nuevos_por_lote, "sha_taxonomia": sha_tax, "sha_marco": sha_marco},
                     ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()

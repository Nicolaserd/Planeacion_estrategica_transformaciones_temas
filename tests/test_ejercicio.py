import sys
from pathlib import Path

import pandas as pd
import pytest
import yaml

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
import ejercicio as ej  # noqa: E402

DATOS = next(iter(sorted((RAIZ / "data" / "raw").glob("*(UC PARA LA VIDA)*.xlsx"))), RAIZ / "data" / "raw" / "no_existe.xlsx")


def test_resolver_transformacion_uc_para_la_vida():
    t = ej.resolver_transformacion("UC PARA LA VIDA")
    assert t["id"] == "T5" and t["nombre_oficial"] == "UC para la Vida"
    assert len(t["componentes"]) == 9
    assert "Felicidad y bienestar" in t["componentes"]


def test_resolver_transformacion_inexistente():
    with pytest.raises(SystemExit):
        ej.resolver_transformacion("UC Inexistente")


def test_anonimiza_encabezados_whatsapp():
    from f03_limpieza_anonimizacion import anonimizar
    conteo = {}
    t = anonimizar("[19/10, 10:06 a. m.] Ana María Pérez Gómez: Conexión a Internet", conteo)
    assert t == "[19/10, 10:06 a. m.] [PERSONA]: Conexión a Internet" and conteo["PERSONA"] == 1
    firma = anonimizar("Con gusto ayudo…Ana María Pérez Gómez, facultad de ciencias", {})
    assert firma == "Con gusto ayudo…[PERSONA], facultad de ciencias"
    assert anonimizar("Más inglés, programa bilingüe", {}) == "Más inglés, programa bilingüe"   # una sola palabra: no es firma


def test_formulario_general_ucundinamarca():
    t = ej.resolver_transformacion("UCundinamarca")
    assert t["general"] and t["nombre_oficial"] == "UCundinamarca"
    assert t["componentes"] == ["UC Inteligente", "UC Translocal", "UC Digital", "UC Emprendedora e Innovadora",
                                "UC para la Vida"]
    assert ej.componente_valido("UC Digital", t) and ej.componente_valido("Felicidad y bienestar", t)
    assert ej.etiquetas_ancla(t)[1] == "Transformación del Plan"
    assert ej.slug_ejercicio("Todas", ["2026-08-31", "2026-10-27"], "UCundinamarca") == \
        "todas_2026-08-31_a_2026-10-27_ucundinamarca"


def test_componente_valido():
    t = ej.resolver_transformacion("UC para la Vida")
    assert ej.componente_valido("Convivencia", t)
    assert ej.componente_valido("condiciones que favorezcan el aprendizaje", t)
    assert ej.componente_valido(ej.FUERA, t)
    assert ej.componente_valido(ej.GENERAL, t)
    assert ej.componente_valido("cultura institucional orientada a la vida", t)
    assert not ej.componente_valido("Inteligencia artificial", t)   # es de T1 y T3, no de T5
    assert not ej.componente_valido("Vida", t)                       # fragmento demasiado corto


def test_minimo_comentarios_y_componentes_del_tema():
    assert ej.minimo_comentarios(17) == 3 and ej.minimo_comentarios(60) == 3
    assert ej.minimo_comentarios(78) == 4 and ej.minimo_comentarios(119) == 6
    t = {"componente": "Convivencia", "otros_componentes": ["Valores democráticos"]}
    assert ej.componentes_del_tema(t) == ["Convivencia", "Valores democráticos"]
    assert ej.componentes_del_tema({"componente": ej.FUERA}) == [ej.FUERA]


def test_copia_anonimizada(tmp_path):
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Correo electrónico", "Nombre", "¿Qué ajustaría?", "Tema principal (Jev)"])
    ws.append([7, "ana@x.co", "Ana Pérez", "Que el profesor Juan explique mejor", "Docentes"])
    ws.append([8, "luis@x.co", "Luis Gómez", None, "(sin comentario)"])
    ruta = tmp_path / "Clasificacion_Prueba.xlsx"
    wb.save(ruta)
    com = pd.DataFrame({"id_resp": [7, 8], "comentario_anon": ["Que el profesor [PERSONA] explique mejor", ""]})
    destino = ej.copia_anonimizada(ruta, "¿Qué ajustaría?", com)
    assert destino.name == "Clasificacion_Prueba_anonimizado.xlsx"
    filas = list(load_workbook(destino).active.iter_rows(min_row=2, values_only=True))
    assert filas[0] == (7, "[CORREO]", "[NOMBRE]", "Que el profesor [PERSONA] explique mejor", "Docentes")
    assert filas[1] == (8, "[CORREO]", "[NOMBRE]", None, "(sin comentario)")
    assert load_workbook(ruta).active.cell(row=2, column=3).value == "Ana Pérez"   # el original no se toca


def test_slug_y_nombres():
    fechas = ["2026-08-19", "2026-08-18"]
    assert ej.slug_ejercicio("Ubaté", fechas, "UC para la Vida") == "ubate_2026-08-18_a_2026-08-19_uc_para_la_vida"
    assert ej.nombre_archivo("Temas", "Ubaté", fechas, "UC para la Vida") == \
        "Temas_Ubate_2026-08-18_a_2026-08-19_UC_para_la_Vida.xlsx"
    assert ej.nombre_hoja("Ubaté", fechas) == "Ubaté 18-19 ago 2026"
    assert len(ej.nombre_hoja("Extensión Zipaquirá", ["2026-08-31", "2026-09-22"])) <= 31
    assert ej.parte_fechas(["2026-09-07"]) == "2026-09-07"
    assert ej.carpeta_transformacion("UC para la Vida").name == "UC para la Vida"


@pytest.mark.skipif(not DATOS.exists(), reason="datos locales no disponibles")
def test_filtro_ubate():
    mapeo = yaml.safe_load((RAIZ / "config" / "mapeo_columnas.yaml").read_text(encoding="utf-8"))
    df = pd.read_excel(DATOS)
    fechas = ["2026-08-18", "2026-08-19"]
    sub, inf = ej.filtrar(df, mapeo, "Ubaté", fechas, "UC para la Vida")
    assert len(sub) == 65
    assert sorted(x["ID"] for x in inf["fuera_de_fecha_no_incluidas"]) == [408, 493]
    sub2, inf2 = ej.filtrar(df, mapeo, "Ubaté", fechas, "UC para la Vida", incluir_ids=[408, 493])
    assert len(sub2) == 67 and not inf2["fuera_de_fecha_no_incluidas"]
    with pytest.raises(SystemExit):
        ej.filtrar(df, mapeo, "Ubaté", fechas, "UC Digital")

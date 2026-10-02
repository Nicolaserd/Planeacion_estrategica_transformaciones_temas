"""Ejecuta un paso del flujo (src/ejercicio.py) para varios ejercicios a la vez.

Uso, desde la raíz del proyecto:
  .venv\\Scripts\\python.exe scripts/lote.py validar    [filtro]
  .venv\\Scripts\\python.exe scripts/lote.py clasificar [filtro]
  .venv\\Scripts\\python.exe scripts/lote.py exportar   [filtro]
  .venv\\Scripts\\python.exe scripts/lote.py verificar  [filtro]

`filtro` es cualquier parte del slug: "ubate", "uc_digital", "chia_2026-09-21_a_2026-09-22_uc_translocal"...
Sin filtro se procesan todos los ejercicios de config/ejercicios/.
"""
import contextlib
import io
import json
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
import ejercicio as ej  # noqa: E402


def ejercicios(filtro: str) -> list[str]:
    fichas = sorted((RAIZ / "config" / "ejercicios").glob("*/ficha.yaml"))
    return [p.parent.name for p in fichas if filtro in p.parent.name]


def verificar(s: str) -> list[str]:
    """Los dos Excel existen, tienen una hoja, una fila por respuesta y una columna por tema."""
    ficha, temas = ej.cargar(s)
    carpeta = RAIZ / "outputs" / ficha["transformacion"]["nombre_oficial"]
    r1, r2 = carpeta / ficha["archivos_salida"]["temas"], carpeta / ficha["archivos_salida"]["clasificacion"]
    if not (r1.exists() and r2.exists()):
        return ["falta algún Excel"]
    problemas = []
    w1, w2 = load_workbook(r1), load_workbook(r2)
    if len(w1.sheetnames) != 1 or len(w2.sheetnames) != 1:
        problemas.append("más de una hoja")
    ws = w2.active
    enc = [ws.cell(row=1, column=j).value for j in range(1, ws.max_column + 1)]
    if ws.max_row - 1 != ficha["N_filas"]:
        problemas.append(f"{ws.max_row - 1} filas en el Excel de clasificación y {ficha['N_filas']} respuestas")
    if "% confianza (Jev)" not in enc:
        problemas.append("falta la columna '% confianza (Jev)'")
    else:
        cols_tema = enc[enc.index("% confianza (Jev)") + 1:]
        if sorted(cols_tema) != sorted(t["nombre"] for t in temas["temas"]):
            problemas.append("las columnas de tema no coinciden con temas.yaml")
    return problemas


def main():
    if len(sys.argv) < 2 or sys.argv[1] not in ("validar", "clasificar", "exportar", "verificar"):
        raise SystemExit(__doc__)
    paso, filtro = sys.argv[1], (sys.argv[2] if len(sys.argv) > 2 else "")
    lista = ejercicios(filtro)
    if not lista:
        raise SystemExit(f"Ningún ejercicio coincide con '{filtro}'.")
    fallos = 0
    for s in lista:
        buf = io.StringIO()
        try:
            with contextlib.redirect_stdout(buf):
                if paso == "validar":
                    r = ej.validar(s, imprimir=False)
                    detalle = f"temas={r['temas']} mínimo={r['minimo_por_tema']} sin_tema(lectura)={r['sin_tema']}"
                elif paso == "clasificar":
                    ej.clasificar(s)
                    detalle = buf.getvalue().strip().splitlines()[-1]
                elif paso == "exportar":
                    ej.exportar(s)
                    r = json.loads(buf.getvalue())
                    detalle = f"filas={r['filas_excel2']} asignaciones={r['asignaciones']} sin_tema={r['sin_tema']}"
                else:
                    p = verificar(s)
                    if p:
                        raise SystemExit("; ".join(p))
                    detalle = "Excel correctos"
            print(f"OK    {s:<62} {detalle}")
        except SystemExit as e:
            fallos += 1
            print(f"FALLA {s}\n      {e}")
    print(f"{len(lista)} ejercicio(s), {fallos} con fallas.")


if __name__ == "__main__":
    main()

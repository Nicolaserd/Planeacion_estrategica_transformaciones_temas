"""Temas con BERT (BERTopic), en paralelo e independiente del flujo de ejercicios.

Parte solo de los datos en crudo (data/raw/*.xlsx, los formularios de Forms). No usa los ejercicios, sus fechas,
los temas de Claude ni las clasificaciones de Jev. Para cada unidad regional × transformación y para cada
transformación (todas las unidades juntas) obtiene:
  - el número óptimo de temas: el mayor k en el que todos los temas se reproducen al remuestrear (estabilidad por
    tema: Jaccard medio con su mejor pareja en 20 submuestras del 80 %; umbrales de Hennig (2007): >= 0,75 estable,
    >= 0,60 moderada); silhouette, coherencia NPMI y diversidad se reportan como descripción; HDBSCAN automático
    como referencia,
  - los temas (palabras clave c-TF-IDF de BERTopic y nombre legible en config/bertopic/<grupo>.yaml),
  - el tema de cada comentario y su afinidad (similitud coseno con el centroide del tema).

Uso, desde la raíz del proyecto:
  .venv\\Scripts\\python.exe scripts/bertopic_temas.py preparar           # crudo -> texto limpio y anonimizado
  .venv\\Scripts\\python.exe scripts/bertopic_temas.py modelar  [filtro]  # vectores BERT, número óptimo y temas
  .venv\\Scripts\\python.exe scripts/bertopic_temas.py mostrar  [filtro]  # temas y comentarios representativos
  .venv\\Scripts\\python.exe scripts/bertopic_temas.py exportar [filtro]  # Excel en outputs_bertopic/

`filtro` es cualquier parte del id del grupo, por ejemplo "uc_digital", "fusagasuga" o "todas_las_unidades".
Privacidad: la tabla de comentarios nunca carga Nombre ni Correo; el texto se anonimiza con
src/f03_limpieza_anonimizacion.py y se revisa con los patrones de scripts/revisar_nombres.py. El modelo BERT
corre en el equipo: ningún comentario sale del computador.
"""
import hashlib
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook
from sklearn.cluster import KMeans
from sklearn.feature_extraction.text import CountVectorizer
from sklearn.metrics import silhouette_score

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
sys.path.insert(0, str(RAIZ / "scripts"))
import excel_estilo as estilo  # noqa: E402
import metricas  # noqa: E402
from f03_limpieza_anonimizacion import anonimizar, motivo_exclusion, normalizar, tokens_nombres  # noqa: E402
from revisar_nombres import FIRMA, SECUENCIA, WHATSAPP, nombres_de_pila  # noqa: E402

# ---------- parámetros del análisis (propios de BERTopic, independientes de las reglas de los ejercicios) ----------
MODELO = "sentence-transformers/paraphrase-multilingual-mpnet-base-v2"
SEMILLA = 42
K_MAX = 15                 # máximo de temas que se prueban
MIN_TEMA_ABS = 5           # un tema necesita 5 comentarios...
MIN_TEMA_FRAC = 0.03       # ...o el 3 % del grupo, lo que sea mayor
N_MIN_GRUPO = 20           # con menos comentarios no se modela la unidad (entra en el modelo de su transformación)
N_SIN_UMAP = 30            # con menos comentarios la referencia HDBSCAN no reduce con UMAP
TOP_PALABRAS = 10
REPETICIONES = 20          # submuestras para medir la estabilidad
FRACCION_SUB = 0.8         # tamaño de cada submuestra
JACCARD_ESTABLE = 0.75     # Hennig (2007, clusterboot): un tema con Jaccard medio >= 0,75 es estable
JACCARD_MODERADA = 0.60    # entre 0,60 y 0,75 indica un patrón; por debajo, el tema no se sostiene (débil)
# Los temas se agrupan sobre los vectores BERT completos (sin UMAP): con UMAP la primera división separaba
# comentarios cortos de largos (estilo, no tema) y la estabilidad dependía del azar de UMAP.

DIR_INTERIM = RAIZ / "data" / "interim_bertopic"
DIR_MODELOS = DIR_INTERIM / "modelos"
CACHE = RAIZ / "cache" / "bertopic_embeddings.npz"
SALIDA = RAIZ / "outputs_bertopic"
DIR_NOMBRES = RAIZ / "config" / "bertopic"
TODAS = "Todas las unidades"
SIN_UNIDAD = "(sin unidad)"

# Formulario de Forms (texto entre paréntesis del nombre del archivo) -> nombre oficial de la transformación
FORMULARIOS = {"UC INTELIGENTE": "UC Inteligente", "UC TRANSLOCAL": "UC Translocal", "UC DIGITAL": "UC Digital",
               "UC EMPRENDEDORA E INNOVADORA": "UC Emprendedora e Innovadora", "UC PARA LA VIDA": "UC para la Vida",
               "UCUNDINAMARCA": "UCundinamarca"}
ORDEN_T = list(FORMULARIOS.values())
PREGUNTAS_ABIERTAS = ("¿Qué ajustaría en esta transformación?", "Comparte aquí tu aporte:")
PERSONALES = {"nombre", "correo electronico", "correo"}

# Palabras vacías del español y de la pregunta (no aportan al tema); se comparan sin tildes
STOP = {normalizar(w) for w in """
a al algo alguien algun alguna algunas alguno algunos ambos ante antes aqui asi aun aunque bajo bastante bien cada casi
como con conmigo contra cual cuales cualquier cuando cuanto cuanta de del demas dentro desde donde dos durante e el ella
ellas ello ellos en entonces entre era eran eres es esa esas ese eso esos esta estaba estaban estado estamos estan estar
estas este esto estos estoy fue fueron fui ha haber habia habian han has hasta hay haya hayan he hemos hoy la las le
les lo los mas me mi mia mis mismo misma mismos mismas mucha muchas mucho muchos muy nada ni ninguna ninguno no nos
nosotros nuestra nuestras nuestro nuestros o otra otras otro otros para pero poca pocas poco pocos por porque pues que
quien quienes se sea sean segun ser si sido siempre sin sino sobre solo son su sus tal tambien tan tanto te tener tenga
tengan tenemos tiene tienen toda todas todo todos tu tus u un una unas uno unos usted ustedes va van vez veces y ya yo
cosa cosas parte etc ademas manera forma hacer hace hacen hacia puede pueden podria podrian poder deberia deberian debe
deben deber debemos seria serian creo considero consideramos pienso pensamos opinamos gustaria quisieramos importante
necesario necesita necesitan mejor mejores mayor menor menos mas lugar tipo cuenta realmente ademas mediante traves
universidad udec ucundinamarca cundinamarca transformacion transformaciones ajustaria ajustariamos ajustar ajuste ajustes
uc respecto ello aspecto aspectos tema temas sentido punto puntos falta hacerlo tener tenido sido estan esten
""".split()}


# ---------------------------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------------------------

def slug(texto: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", normalizar(str(texto))).strip("_")


def limpio(texto: str) -> str:
    """Para nombres de archivo: sin tildes ni espacios, conservando mayúsculas ("UC para la Vida" -> "UC_para_la_Vida")."""
    sin_tildes = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "_", sin_tildes).strip("_")


def tamano_minimo(n: int) -> int:
    return max(MIN_TEMA_ABS, math.ceil(MIN_TEMA_FRAC * n))


def formulario_de_archivo(ruta: Path) -> str | None:
    """Transformación del formulario según el texto entre paréntesis del nombre del archivo de Forms."""
    for parte in re.findall(r"\(([^)]*)\)", ruta.name.replace("\xa0", " ")):
        clave = re.sub(r"\s+", " ", parte).strip().upper()
        if clave in FORMULARIOS:
            return FORMULARIOS[clave]
    return None


def columnas_a_leer(columnas, mapeo: dict) -> list[str]:
    """Columnas que se cargan del crudo: nunca Nombre ni Correo."""
    queridas = {mapeo["id_resp"], mapeo["fecha"], mapeo["unidad_regional"], mapeo["grupo"], mapeo["programa"],
                *PREGUNTAS_ABIERTAS}
    return [c for c in columnas if c in queridas and normalizar(str(c)) not in PERSONALES]


def texto_para_palabras(t: str) -> str:
    """Texto para las palabras clave: sin las marcas de anonimización ([PERSONA], [URL]...)."""
    return re.sub(r"\[[A-Z_]+\]", " ", str(t))


def analizador(ngramas=(1, 2)):
    """Tokens de 3 o más letras, sin palabras vacías (comparadas sin tildes)."""
    base = CountVectorizer(token_pattern=r"(?u)\b[^\W\d_]{3,}\b", lowercase=True, ngram_range=(1, 1)).build_analyzer()

    def analizar(doc):
        toks = [w for w in base(doc) if normalizar(w) not in STOP]
        out = list(toks)
        if ngramas[1] >= 2:
            out += [f"{a} {b}" for a, b in zip(toks, toks[1:])]
        return out
    return analizar


def vectorizador(n_docs: int, ngramas=(1, 2)) -> CountVectorizer:
    return CountVectorizer(analyzer=analizador(ngramas), min_df=2 if n_docs >= 100 else 1)


# ---------------------------------------------------------------------------------------------
# Métricas para el número de temas
# ---------------------------------------------------------------------------------------------

def palabras_ctfidf(X, etiquetas: np.ndarray, k: int, vocab: np.ndarray, top: int = TOP_PALABRAS) -> list[list[int]]:
    """Índices de las `top` palabras de cada tema según c-TF-IDF (la fórmula de BERTopic)."""
    C = np.vstack([np.asarray(X[etiquetas == t].sum(axis=0)).ravel() for t in range(k)])
    tf = C / np.maximum(C.sum(axis=1, keepdims=True), 1)
    A = C.sum() / k
    idf = np.log(1 + A / np.maximum(C.sum(axis=0), 1))
    w = tf * idf
    return [[int(j) for j in np.argsort(-w[t])[:top] if C[t, j] > 0] for t in range(k)]


def npmi(temas: list[list[int]], D: np.ndarray) -> float:
    """Coherencia NPMI media de las palabras de cada tema (coocurrencia en los mismos comentarios)."""
    n = D.shape[0]
    valores = []
    for idx in temas:
        for a in range(len(idx)):
            for b in range(a + 1, len(idx)):
                i, j = idx[a], idx[b]
                pij = np.logical_and(D[:, i], D[:, j]).sum() / n
                if pij == 0:
                    valores.append(-1.0)
                elif pij == 1:
                    valores.append(1.0)
                else:
                    pi, pj = D[:, i].sum() / n, D[:, j].sum() / n
                    valores.append(math.log(pij / (pi * pj)) / -math.log(pij))
    return float(np.mean(valores)) if valores else 0.0


def diversidad(temas: list[list[int]]) -> float:
    todas = [i for t in temas for i in t]
    return len(set(todas)) / len(todas) if todas else 0.0


def agrupar(E: np.ndarray, k: int) -> np.ndarray:
    return KMeans(n_clusters=k, n_init=20, random_state=SEMILLA).fit_predict(E)


def estabilidad_por_tema(E: np.ndarray, etiquetas: np.ndarray, k: int) -> np.ndarray:
    """Jaccard medio de cada tema con su mejor pareja al reagrupar submuestras del 80 % (Hennig 2007)."""
    rng = np.random.default_rng(SEMILLA)
    J = np.zeros((REPETICIONES, k))
    for r in range(REPETICIONES):
        idx = rng.choice(len(E), int(FRACCION_SUB * len(E)), replace=False)
        sub = KMeans(n_clusters=k, n_init=10, random_state=r).fit_predict(E[idx])
        orig = etiquetas[idx]
        for t in range(k):
            a = orig == t
            if a.any():
                J[r, t] = max(np.logical_and(a, sub == s).sum() / np.logical_or(a, sub == s).sum() for s in range(k))
    return J.mean(axis=0)


def nivel_estabilidad(jaccard: float) -> str:
    return "estable" if jaccard >= JACCARD_ESTABLE else "moderada" if jaccard >= JACCARD_MODERADA else "débil"


def curva_k(E: np.ndarray, docs: list[str], minimo: int, k_max: int = K_MAX) -> list[dict]:
    """Prueba k = 2..k_max con KMeans sobre los vectores: estabilidad (criterio), silhouette, NPMI y diversidad."""
    vec = vectorizador(len(docs), ngramas=(1, 1))
    X = vec.fit_transform(docs)
    D = (X > 0).toarray()
    filas = []
    for k in range(2, min(k_max, len(E) // minimo) + 1):
        lab = agrupar(E, k)
        tam = int(np.bincount(lab, minlength=k).min())
        top = palabras_ctfidf(X, lab, k, vec.get_feature_names_out())
        j = estabilidad_por_tema(E, lab, k)
        filas.append(dict(k=k, tam_min=tam, cumple=tam >= minimo, estabilidad=float(j.min()),
                          estabilidad_media=float(j.mean()), estabilidad_temas=[float(x) for x in j],
                          silhouette=float(silhouette_score(E, lab, metric="cosine")),
                          npmi=npmi(top, D), diversidad=diversidad(top)))
    elegir_k(filas)
    return filas


def elegir_k(filas: list[dict]) -> int | None:
    """El mayor k que cumple el tamaño mínimo y en el que todos los temas son estables (Jaccard del tema más débil
    >= 0,75); si ninguno, el mayor con estabilidad moderada (>= 0,60); si ninguno, el más estable (débil)."""
    validas = [f for f in filas if f["cumple"]]
    for f in filas:
        f["nivel"], f["elegido"] = nivel_estabilidad(f["estabilidad"]), False
    if not validas:
        return None
    estables = [f for f in validas if f["estabilidad"] >= JACCARD_ESTABLE]
    moderadas = [f for f in validas if f["estabilidad"] >= JACCARD_MODERADA]
    if estables or moderadas:
        mejor = max(estables or moderadas, key=lambda f: f["k"])
    else:
        mejor = max(validas, key=lambda f: (round(f["estabilidad"], 9), -f["k"]))
    mejor["elegido"] = True
    return mejor["k"]


# ---------------------------------------------------------------------------------------------
# preparar: crudo -> comentarios limpios y anonimizados
# ---------------------------------------------------------------------------------------------

def preparar():
    mapeo = yaml.safe_load((RAIZ / "config" / "mapeo_columnas.yaml").read_text(encoding="utf-8"))
    filas, resumen, conteo = [], [], {}
    for ruta in sorted((RAIZ / "data" / "raw").glob("*.xlsx")):
        T = formulario_de_archivo(ruta)
        if T is None:
            print(f"  (se omite {ruta.name}: no es un formulario conocido)")
            continue
        cols = columnas_a_leer(pd.read_excel(ruta, nrows=0).columns, mapeo)
        df = pd.read_excel(ruta, usecols=cols)
        col_txt = next(c for c in PREGUNTAS_ABIERTAS if c in df.columns)
        col_nombre = next((c for c in pd.read_excel(ruta, nrows=0).columns if normalizar(str(c)) == "nombre"), None)
        toks = tokens_nombres(pd.read_excel(ruta, usecols=[col_nombre])) if col_nombre else set()
        df["unidad"] = df[mapeo["unidad_regional"]].fillna(SIN_UNIDAD).astype(str).str.strip()
        for U, sub in df.groupby("unidad"):
            motivos = sub[col_txt].map(motivo_exclusion)
            resumen.append(dict(transformacion=T, unidad=U, respuestas=len(sub), con_contenido=int(motivos.isna().sum()),
                                **{f"excluidas_{m}": int((motivos == m).sum()) for m in motivos.dropna().unique()}))
        for _, x in df.iterrows():
            if motivo_exclusion(x[col_txt]) is not None:
                continue
            anon = anonimizar(str(x[col_txt]).strip(), conteo)
            filas.append(dict(transformacion=T, unidad=x["unidad"], id_resp=int(x[mapeo["id_resp"]]),
                              id_com=f"{slug(T)}-{int(x[mapeo['id_resp']])}",
                              fecha=pd.Timestamp(x[mapeo["fecha"]]).strftime("%Y-%m-%d"),
                              rol=x.get(mapeo["grupo"]), programa=x.get(mapeo["programa"]),
                              comentario_anon=anon,
                              token_nombre=int(any(normalizar(w) in toks for w in re.findall(r"[^\W\d_]{4,}", anon)))))
    com = pd.DataFrame(filas)
    # control de nombres que la anonimización no reconoció (solo se imprime el id)
    pilas = nombres_de_pila()
    hallazgos = []
    for _, x in com.iterrows():
        t = x.comentario_anon
        tipos = [n for n, ok in (("whatsapp", WHATSAPP.search(t)), ("firma", FIRMA.search(t)),
                                 ("nombre_probable", any(normalizar(m.group(1)) in pilas for m in SECUENCIA.finditer(t))))
                 if ok]
        if tipos:
            hallazgos.append(f"{x.id_com} {tipos}")
    DIR_INTERIM.mkdir(parents=True, exist_ok=True)
    com.to_parquet(DIR_INTERIM / "comentarios.parquet", index=False)
    pd.DataFrame(resumen).fillna(0).to_parquet(DIR_INTERIM / "resumen_crudo.parquet", index=False)
    (DIR_INTERIM / "preparar.json").write_text(json.dumps({"anonimizacion": conteo, "revision_nombres": len(hallazgos)},
                                                          ensure_ascii=False, indent=1), encoding="utf-8")
    tabla = pd.DataFrame(resumen).pivot_table(index="unidad", columns="transformacion", values="con_contenido",
                                              aggfunc="sum", fill_value=0)
    print(tabla[[t for t in ORDEN_T if t in tabla.columns]].to_string())
    print(f"\n{len(com)} comentarios con contenido · anonimización: {conteo}")
    print(f"Revisión de nombres: {len(hallazgos)} comentario(s) a revisar")
    for h in hallazgos:
        print("  ", h)
    if hallazgos:
        raise SystemExit("Hay nombres sin anonimizar: ampliar src/f03_limpieza_anonimizacion.py antes de modelar.")


# ---------------------------------------------------------------------------------------------
# modelar: vectores BERT, número óptimo y modelo final
# ---------------------------------------------------------------------------------------------

def _clave(texto: str) -> str:
    return hashlib.sha256(f"{MODELO}\x00{texto}".encode("utf-8")).hexdigest()


_MODELO_ST = None


def modelo_st():
    global _MODELO_ST
    if _MODELO_ST is None:
        from sentence_transformers import SentenceTransformer
        _MODELO_ST = SentenceTransformer(MODELO, device="cpu")
    return _MODELO_ST


def vectores(textos: list[str]) -> np.ndarray:
    """Vectores BERT normalizados, con caché en cache/bertopic_embeddings.npz."""
    guardado = {}
    if CACHE.exists():
        z = np.load(CACHE, allow_pickle=False)
        guardado = dict(zip(z["claves"], z["vectores"]))
    faltan = sorted({t for t in textos if _clave(t) not in guardado})
    if faltan:
        print(f"  calculando {len(faltan)} vectores con {MODELO}...")
        nuevos = modelo_st().encode(faltan, batch_size=32, normalize_embeddings=True, show_progress_bar=True)
        guardado.update({_clave(t): v for t, v in zip(faltan, nuevos)})
        CACHE.parent.mkdir(parents=True, exist_ok=True)
        claves = np.array(list(guardado))
        np.savez(CACHE, claves=claves, vectores=np.vstack([guardado[c] for c in claves]))
    return np.vstack([guardado[_clave(t)] for t in textos]).astype(np.float32)


def grupos(com: pd.DataFrame) -> list[dict]:
    out = []
    for T in ORDEN_T:
        sub = com[com.transformacion == T]
        if sub.empty:
            continue
        out.append(dict(id=f"{slug(T)}__todas_las_unidades", nivel="transformacion", T=T, U=TODAS, idx=sub.index))
        for U in sorted(u for u in sub.unidad.unique() if u != SIN_UNIDAD):
            out.append(dict(id=f"{slug(T)}__{slug(U)}", nivel="unidad", T=T, U=U, idx=sub[sub.unidad == U].index))
    return out


def referencia_hdbscan(E: np.ndarray, minimo: int) -> dict:
    """BERTopic en su modo automático (UMAP + HDBSCAN): cuántos temas encuentra y cuántos comentarios deja sin tema."""
    from hdbscan import HDBSCAN
    R = E
    if len(E) >= N_SIN_UMAP:
        from umap import UMAP
        R = UMAP(n_neighbors=min(15, len(E) - 1), n_components=5, min_dist=0.0, metric="cosine",
                 random_state=SEMILLA).fit_transform(E)
    lab = HDBSCAN(min_cluster_size=minimo, metric="euclidean", cluster_selection_method="eom").fit_predict(R)
    return dict(k=int(len(set(lab) - {-1})), sin_tema=float(np.mean(lab == -1)))


def modelar_grupo(g: dict, com: pd.DataFrame) -> dict:
    from bertopic import BERTopic
    from bertopic.dimensionality import BaseDimensionalityReduction
    from bertopic.representation import MaximalMarginalRelevance

    sub = com.loc[g["idx"]]
    textos = list(sub.comentario_anon)
    docs = [texto_para_palabras(t) for t in textos]
    E = vectores(textos)
    N, minimo = len(sub), tamano_minimo(len(sub))
    info = dict(id=g["id"], nivel=g["nivel"], T=g["T"], U=g["U"], N=N, minimo=minimo, modelo=MODELO,
                fechas=[sub.fecha.min(), sub.fecha.max()])
    if g["nivel"] == "unidad" and N < N_MIN_GRUPO:
        return info | dict(modelado=False, motivo=f"menos de {N_MIN_GRUPO} comentarios")
    curva = curva_k(E, docs, minimo)
    k = next((f["k"] for f in curva if f["elegido"]), None)
    info |= dict(curva=curva, hdbscan=referencia_hdbscan(E, minimo))
    if k is None:
        return info | dict(modelado=False, motivo="ningún k deja temas del tamaño mínimo")
    elegido = next(f for f in curva if f["elegido"])
    info |= dict(estabilidad=elegido["estabilidad"], nivel_estabilidad=elegido["nivel"])
    modelo = BERTopic(embedding_model=modelo_st(), umap_model=BaseDimensionalityReduction(),
                      hdbscan_model=KMeans(n_clusters=k, n_init=20, random_state=SEMILLA),
                      vectorizer_model=vectorizador(N), representation_model=MaximalMarginalRelevance(diversity=0.3),
                      top_n_words=TOP_PALABRAS, calculate_probabilities=False)
    etiquetas = np.array(modelo.fit_transform(docs, embeddings=E)[0])
    lab_km = agrupar(E, k)   # misma partición que en la curva (BERTopic solo renumera los temas por tamaño)
    temas, afinidad = [], np.zeros(N)
    for t in sorted(set(etiquetas)):
        miembros = np.where(etiquetas == t)[0]
        centro = E[miembros].mean(axis=0)
        centro /= np.linalg.norm(centro) or 1
        sim = E[miembros] @ centro
        afinidad[miembros] = sim
        orden = miembros[np.argsort(-sim)]
        estab = elegido["estabilidad_temas"][int(np.bincount(lab_km[miembros]).argmax())]
        temas.append(dict(tema=int(t), n=int(len(miembros)), palabras=[w for w, _ in modelo.get_topic(t) if w.strip()][:8],
                          afinidad_media=float(sim.mean()), estabilidad=estab, nivel=nivel_estabilidad(estab),
                          representativos=[sub.iloc[i].id_com for i in orden[:3]]))
    DIR_MODELOS.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(dict(id_com=sub.id_com.values, tema=etiquetas, afinidad=afinidad)).to_parquet(
        DIR_MODELOS / f"{g['id']}_asignaciones.parquet", index=False)
    return info | dict(modelado=True, k=k, temas=temas)


def modelar(filtro: str):
    com = pd.read_parquet(DIR_INTERIM / "comentarios.parquet")
    DIR_MODELOS.mkdir(parents=True, exist_ok=True)
    for g in grupos(com):
        if filtro not in g["id"]:
            continue
        info = modelar_grupo(g, com)
        (DIR_MODELOS / f"{g['id']}.json").write_text(json.dumps(info, ensure_ascii=False, indent=1), encoding="utf-8")
        if info["modelado"]:
            h = info["hdbscan"]
            print(f"OK  {g['id']:<58} N={info['N']:<4} k*={info['k']:<3} estabilidad={info['estabilidad']:.2f} "
                  f"({info['nivel_estabilidad']:<8}) HDBSCAN={h['k']} ({100 * h['sin_tema']:.0f} % sin tema)")
        else:
            print(f"--  {g['id']:<58} N={info['N']:<4} sin modelar: {info['motivo']}")


# ---------------------------------------------------------------------------------------------
# mostrar: temas y comentarios representativos (para redactar los nombres legibles)
# ---------------------------------------------------------------------------------------------

def firma(palabras: list[str]) -> str:
    return ", ".join(palabras[:5])


def mostrar(filtro: str):
    com = pd.read_parquet(DIR_INTERIM / "comentarios.parquet").set_index("id_com")
    for ruta in sorted(DIR_MODELOS.glob("*.json")):
        info = json.loads(ruta.read_text(encoding="utf-8"))
        if filtro not in info["id"] or not info.get("modelado"):
            continue
        print(f"\n=== {info['id']} · N={info['N']} · k*={info['k']} · estabilidad {info['estabilidad']:.2f} "
              f"({info['nivel_estabilidad']})")
        for t in info["temas"]:
            print(f"  [{t['tema']}] n={t['n']} · {firma(t['palabras'])} | {', '.join(t['palabras'][5:])}")
            for c in t["representativos"]:
                x = com.loc[c]
                marca = " <<token de nombre>>" if x.token_nombre else ""
                print(f"      - {x.comentario_anon[:230]}{marca}")


# ---------------------------------------------------------------------------------------------
# exportar: Excel en outputs_bertopic/
# ---------------------------------------------------------------------------------------------

def nombres_legibles(info: dict) -> dict[int, str]:
    """Nombres de config/bertopic/<grupo>.yaml; si la firma cambió, se avisa y se usan las palabras clave."""
    ruta = DIR_NOMBRES / f"{info['id']}.yaml"
    guardados = {}
    if ruta.exists():
        for t in (yaml.safe_load(ruta.read_text(encoding="utf-8")) or {}).get("temas", []):
            guardados[int(t["tema"])] = t
    out = {}
    for t in info["temas"]:
        g = guardados.get(t["tema"])
        if g and g.get("firma") == firma(t["palabras"]):
            out[t["tema"]] = g["nombre"]
        else:
            if g:
                print(f"  AVISO {info['id']} tema {t['tema']}: cambió la firma; se usan las palabras clave")
            out[t["tema"]] = f"Tema {t['tema'] + 1}: {', '.join(t['palabras'][:3])}"
    return out


def _hoja_temas(ws, info, nombres, com):
    h = info["hdbscan"]
    aviso = {"estable": "Todos los temas se reproducen al remuestrear (estables).",
             "moderada": "Los temas se reproducen en parte al remuestrear (estabilidad moderada): leerlos como patrones.",
             "débil": "ATENCIÓN: ningún número de temas se reproduce al remuestrear (estabilidad débil): los temas "
                      "son orientativos."}[info["nivel_estabilidad"]]
    estilo.titulo(ws, [
        f"BERTopic · {info['T']} · {info['U']}",
        f"{info['N']} comentarios con contenido · fechas {info['fechas'][0]} a {info['fechas'][1]} · modelo {MODELO}",
        f"Número óptimo de temas: {info['k']}, el mayor número de temas que se reproducen al remuestrear (estabilidad "
        f"del tema más débil = {info['estabilidad']:.2f}, {info['nivel_estabilidad']}; temas de {info['minimo']} "
        "comentarios o más). "
        f"{aviso} BERTopic automático (HDBSCAN), como referencia: {h['k']} temas y {100 * h['sin_tema']:.0f} % de "
        "comentarios sin tema. Cada comentario está en un solo tema; % sobre los comentarios con contenido; cada "
        "comentario es de un grupo de 5 a 8 personas.",
    ], 10)
    filas = []
    for t in info["temas"]:
        lo, hi = metricas.wilson(t["n"], info["N"])
        filas.append([t["tema"] + 1, nombres[t["tema"]], ", ".join(t["palabras"]), t["n"], t["n"] / info["N"], lo, hi,
                      t["estabilidad"], t["afinidad_media"], com.loc[t["representativos"][0], "comentario_anon"]])
    estilo.tabla(ws, 5, ["Tema", "Nombre del tema", "Palabras clave (BERTopic)", "Comentarios", "% de comentarios",
                         "IC 95 % inferior", "IC 95 % superior", "Estabilidad del tema (0-1)", "Afinidad media",
                         "Comentario más representativo"],
                 filas, [7, 34, 46, 12, 12, 11, 11, 12, 11, 70], pct_cols=(5, 6, 7), wrap_cols=(2, 3, 10),
                 formatos={8: "0.00", 9: "0.00"})


def _hoja_clasificacion(ws, info, nombres, com, asig):
    d = asig.merge(com.reset_index(), on="id_com").sort_values(["tema", "afinidad"], ascending=[True, False])
    unidad = info["nivel"] == "transformacion"
    enc = ["ID", "Fecha"] + (["Unidad Regional"] if unidad else []) + [
        "Tipo de actor", "Programa", "Comentario (anonimizado)", "Tema", "Nombre del tema", "Afinidad"]
    filas = [[int(x.id_resp), x.fecha] + ([x.unidad] if unidad else []) +
             [x.rol, x.programa, x.comentario_anon, int(x.tema) + 1, nombres[int(x.tema)], float(x.afinidad)]
             for x in d.itertuples()]
    anchos = [8, 11] + ([20] if unidad else []) + [26, 30, 80, 7, 34, 10]
    estilo.tabla(ws, 1, enc, filas, anchos, wrap_cols=(len(enc) - 3,), formatos={len(enc): "0.00"})


def _hoja_numero(ws, info):
    filas = [[f["k"], f["tam_min"], "Sí" if f["cumple"] else "No", f["estabilidad"], f["estabilidad_media"],
              f["nivel"], f["silhouette"], f["npmi"], f["diversidad"], "◀ elegido" if f["elegido"] else ""]
             for f in info["curva"]]
    ultima = estilo.tabla(ws, 1, ["Número de temas (k)", "Tema más chico (comentarios)",
                                  f"¿Cumple el mínimo ({info['minimo']})?", "Estabilidad del tema más débil",
                                  "Estabilidad media de los temas", "Nivel de estabilidad", "Silhouette (coseno)",
                                  "Coherencia NPMI", "Diversidad de palabras", "Elegido"],
                          filas, [14, 16, 16, 14, 14, 13, 14, 14, 14, 12],
                          formatos={4: "0.00", 5: "0.00", 7: "0.000", 8: "0.000", 9: "0.00"})
    h = info["hdbscan"]
    notas = [
        "Criterio: el mayor número de temas que cumple el tamaño mínimo y en el que todos los temas son estables "
        f"(estabilidad del tema más débil >= {JACCARD_ESTABLE:.2f}); si ninguno, el mayor con estabilidad moderada "
        f"(>= {JACCARD_MODERADA:.2f}); si ninguno, el más estable (débil: temas orientativos).",
        f"Estabilidad de un tema: se reagrupan {REPETICIONES} submuestras del {FRACCION_SUB:.0%} de los comentarios y se "
        "mide cuánto coincide el tema con su mejor pareja (Jaccard medio: 1 = siempre el mismo tema; 0 = desaparece). "
        "Umbrales de Hennig (2007): 0,75 estable; 0,60 patrón.",
        "Silhouette: qué tan separados están los temas (de -1 a 1). NPMI: qué tanto aparecen juntas las palabras de cada "
        "tema (de -1 a 1). Diversidad: proporción de palabras clave que no se repiten entre temas. Son descriptivas.",
        f"Referencia BERTopic automático (UMAP + HDBSCAN): {h['k']} temas y {100 * h['sin_tema']:.0f} % de comentarios "
        "sin tema (outliers).",
    ]
    for i, n in enumerate(notas):
        ws.cell(row=ultima + 2 + i, column=1, value=n)


def _hoja_por_unidad(ws, info, nombres, com, asig):
    d = asig.merge(com.reset_index(), on="id_com")
    unidades = sorted(u for u in d.unidad.unique() if u != SIN_UNIDAD) + ([SIN_UNIDAD] if SIN_UNIDAD in set(d.unidad) else [])
    tot = d.groupby("unidad").size()
    filas = []
    for t in info["temas"]:
        en_t = d[d.tema == t["tema"]].groupby("unidad").size()
        filas.append([t["tema"] + 1, nombres[t["tema"]]] + [en_t.get(u, 0) / tot[u] for u in unidades] + [t["n"] / info["N"]])
    enc = ["Tema", "Nombre del tema"] + [f"{u} (n={tot[u]})" for u in unidades] + [f"Total (n={info['N']})"]
    estilo.tabla(ws, 1, enc, filas, [7, 34] + [14] * (len(unidades) + 1),
                 pct_cols=tuple(range(3, len(enc) + 1)), wrap_cols=(2,))
    ws.cell(row=len(filas) + 3, column=1, value="% de los comentarios de cada unidad que quedaron en cada tema "
                                                "(cada columna suma 100 %).")


def exportar(filtro: str):
    com = pd.read_parquet(DIR_INTERIM / "comentarios.parquet").set_index("id_com")
    crudo = pd.read_parquet(DIR_INTERIM / "resumen_crudo.parquet")
    resumen, lista = [], []
    for T in ORDEN_T:
        for ruta in sorted(DIR_MODELOS.glob(f"{slug(T)}__*.json"), key=lambda p: ("todas" not in p.name, p.name)):
            info = json.loads(ruta.read_text(encoding="utf-8"))
            if filtro not in info["id"]:
                continue
            c = crudo[crudo.transformacion == T] if info["U"] == TODAS else crudo[(crudo.transformacion == T) & (crudo.unidad == info["U"])]
            fila = dict(Nivel="Transformación (todas las unidades)" if info["nivel"] == "transformacion" else "Unidad × transformación",
                        Transformación=T, Unidad=info["U"], Respuestas=int(c.respuestas.sum()), Comentarios=info["N"],
                        Modelado="Sí" if info["modelado"] else f"No: {info['motivo']}")
            if info["modelado"]:
                elegido = next(f for f in info["curva"] if f["elegido"])
                carpeta = SALIDA / T
                carpeta.mkdir(parents=True, exist_ok=True)
                nombre = (f"BERTopic_{limpio(T)}_todas_las_unidades.xlsx" if info["nivel"] == "transformacion"
                          else f"BERTopic_{limpio(info['U'])}_{limpio(T)}.xlsx")
                nombres = nombres_legibles(info)
                asig = pd.read_parquet(DIR_MODELOS / f"{info['id']}_asignaciones.parquet")
                assert len(asig) == info["N"] and sum(t["n"] for t in info["temas"]) == info["N"], info["id"]
                wb = Workbook()
                ws = wb.active
                ws.title = "Temas"
                _hoja_temas(ws, info, nombres, com)
                _hoja_clasificacion(wb.create_sheet("Clasificación"), info, nombres, com, asig)
                _hoja_numero(wb.create_sheet("Número de temas"), info)
                if info["nivel"] == "transformacion":
                    _hoja_por_unidad(wb.create_sheet("Por unidad"), info, nombres, com, asig)
                wb.save(carpeta / nombre)
                for t in info["temas"]:
                    lo, hi = metricas.wilson(t["n"], info["N"])
                    lista.append([T, info["U"], info["N"], info["k"], info["nivel_estabilidad"], t["tema"] + 1,
                                  nombres[t["tema"]],
                                  ", ".join(t["palabras"]), t["n"], t["n"] / info["N"], lo, hi, t["estabilidad"],
                                  com.loc[t["representativos"][0], "comentario_anon"]])
                fila |= {"Tamaño mínimo de tema": info["minimo"], "Número óptimo de temas": info["k"],
                         "Estabilidad (tema más débil)": elegido["estabilidad"], "Nivel de estabilidad": elegido["nivel"],
                         "Silhouette": elegido["silhouette"], "Coherencia NPMI": elegido["npmi"],
                         "Diversidad": elegido["diversidad"], "Temas HDBSCAN (referencia)": info["hdbscan"]["k"],
                         "% sin tema HDBSCAN": info["hdbscan"]["sin_tema"], "Archivo": f"{T}/{nombre}"}
                print(f"OK  {T}/{nombre}")
            resumen.append(fila)
    if not filtro:
        cols = ["Nivel", "Transformación", "Unidad", "Respuestas", "Comentarios", "Modelado", "Tamaño mínimo de tema",
                "Número óptimo de temas", "Estabilidad (tema más débil)", "Nivel de estabilidad", "Silhouette", "Coherencia NPMI",
                "Diversidad", "Temas HDBSCAN (referencia)", "% sin tema HDBSCAN", "Archivo"]
        df = pd.DataFrame(resumen).reindex(columns=cols)
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumen"
        estilo.titulo(ws, ["BERTopic · número óptimo de temas por unidad regional y transformación",
                           f"Datos en crudo de los formularios (todas las fechas) · modelo {MODELO} · análisis paralelo, "
                           "independiente de la clasificación con Jev",
                           f"Número óptimo: el mayor número de temas (2 a {K_MAX}) en el que todos los temas se "
                           f"reproducen al remuestrear (estabilidad del tema más débil >= {JACCARD_ESTABLE:.2f} estable; "
                           f">= {JACCARD_MODERADA:.2f} moderada; si no, débil y orientativo), con temas de "
                           f"{MIN_TEMA_ABS} comentarios o el {MIN_TEMA_FRAC:.0%} del grupo. "
                           f"Las unidades con menos de {N_MIN_GRUPO} comentarios no se modelan solas (entran en el "
                           "modelo de su transformación)."],
                      len(cols))
        filas = [[None if (isinstance(v, float) and np.isnan(v)) else v for v in r] for r in df.itertuples(index=False)]
        estilo.tabla(ws, 5, cols, filas, [30, 26, 22, 11, 12, 30, 12, 12, 11, 12, 11, 11, 11, 13, 12, 60],
                     pct_cols=(15,), wrap_cols=(6,), formatos={9: "0.00", 11: "0.000", 12: "0.000", 13: "0.00"})
        SALIDA.mkdir(parents=True, exist_ok=True)
        wb.save(SALIDA / "Resumen_BERTopic.xlsx")
        print(f"OK  Resumen_BERTopic.xlsx ({len(df)} grupos)")
        # Lista única de temas: una fila por tema, filtrable por transformación y unidad regional
        wb = Workbook()
        ws = wb.active
        ws.title = "Temas"
        enc = ["Transformación", "Unidad Regional", "Comentarios del grupo", "Número óptimo de temas",
               "Estabilidad del modelo", "Tema", "Nombre del tema", "Palabras clave (BERTopic)", "Comentarios",
               "% de comentarios", "IC 95 % inferior", "IC 95 % superior", "Estabilidad del tema (0-1)",
               "Comentario más representativo (anonimizado)"]
        estilo.titulo(ws, ["BERTopic · temas por transformación y unidad regional",
                           f"Datos en crudo de los formularios (todas las fechas) · modelo {MODELO} · análisis paralelo, "
                           "independiente de la clasificación con Jev",
                           "«Todas las unidades» es el modelo de la transformación con todas las unidades juntas. Cada "
                           "comentario está en un solo tema; % sobre los comentarios del grupo. Estabilidad: >= 0,75 "
                           "estable; 0,60 a 0,75 moderada; < 0,60 débil (temas orientativos)."], len(enc))
        estilo.tabla(ws, 5, enc, lista, [26, 22, 12, 11, 12, 7, 36, 46, 11, 11, 10, 10, 12, 70],
                     pct_cols=(10, 11, 12), wrap_cols=(7, 8, 14), formatos={13: "0.00"})
        wb.save(SALIDA / "Lista_temas_BERTopic.xlsx")
        print(f"OK  Lista_temas_BERTopic.xlsx ({len(lista)} temas)")


def main():
    if len(sys.argv) < 2 or sys.argv[1] not in ("preparar", "modelar", "mostrar", "exportar"):
        raise SystemExit(__doc__)
    paso, filtro = sys.argv[1], (sys.argv[2] if len(sys.argv) > 2 else "")
    {"preparar": lambda: preparar(), "modelar": lambda: modelar(filtro), "mostrar": lambda: mostrar(filtro),
     "exportar": lambda: exportar(filtro)}[paso]()


if __name__ == "__main__":
    main()
